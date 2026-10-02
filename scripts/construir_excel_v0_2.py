# -*- coding: utf-8 -*-
"""Constrói excel/lobo_insights_industrial_v0_2.xlsx (V0.2 - Analytics & Dashboard).

Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio.

Princípio: a base da V0.1.1 está CONGELADA. Este módulo não altera `construir_excel.py`:
ele IMPORTA `construir()` (que gera as abas de dados, Qualidade, Indicadores e Parametros
exatamente como na V0.1.1) e acrescenta camadas por cima:

    tbConsumo / tbEstoque / tbMateriais      (dados)
      -> Indicadores                          (6 KPIs oficiais e resumos - V0.1.1, inalterados)
        -> Analises                           (cálculo: rankings, Pareto, concentração, fila de estoque, conferências)
          -> DASHBOARD                        (visualização: cartões, gráficos, leituras)

Tudo o que aparece no DASHBOARD é fórmula (nenhum número digitado). Não há previsão.
O recálculo/cache usa as mesmas funções da V0.1.1 (LibreOffice apenas para gravar os valores calculados).
"""
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.drawing.text import Font as DFont
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.xml.functions import Element

from construir_excel import (AVISO, LibreOfficeIndisponivel, construir, injetar_cache, recalcular,
                             sha256)
from lobo_common import DATA, EXCEL

NOME_XLSX_V02 = "lobo_insights_industrial_v0_2.xlsx"
VERSAO = "V0.2"
S_DASH, S_ANA = "DASHBOARD", "Analises"

# ------------------------------------------------------------------ identidade visual
FONTE = "Arial"
NAVY, NAVY_D = "1F3A5F", "12263F"        # azul-marinho (herdado da V0.1) e variante escura do cabeçalho
COBRE = "C9772B"                          # acento industrial (custo)
AZUL = "2F6690"                           # série de unidades
GRAFITE, CINZA_T = "37474F", "5B6B7A"
BG, BRANCO, LINHA = "F3F5F8", "FFFFFF", "D5DBE1"
CINZA_M = "9AA5B1"
VERM, VERM_E = "B03A2E", "8E1B1B"
COR_AZUL_CLARO, COR_COBRE_CLARO = "B7C9DA", "E6CBB0"

fill = lambda cor: PatternFill("solid", fgColor=cor, bgColor=cor)      # noqa: E731
lado = lambda cor, est="thin": Side(style=est, color=cor)             # noqa: E731
FILL_REPOR = PatternFill("solid", bgColor="F8CBCB", fgColor="F8CBCB")
FILL_ATEN = PatternFill("solid", bgColor="FFE9A8", fgColor="FFE9A8")
FILL_OK = PatternFill("solid", bgColor="CDEBD6", fgColor="CDEBD6")

f = lambda sz=10, b=False, cor="222222", i=False, u=None: Font(name=FONTE, size=sz, bold=b, italic=i, color=cor, underline=u)  # noqa: E731

# ------------------------------------------------------------------ geometria do DASHBOARD
COL_INI, COL_FIM = 2, 13                  # colunas B..M (12 colunas de mesma largura)
LARG_COL = 13
METADE = 6                                # cada painel duplo ocupa 6 colunas
KPI_LABEL, KPI_VALOR, KPI_LEIT = 6, 7, 8  # linhas dos cartões (fixas: as conferências apontam para elas)
KPI_COLS = [2, 4, 6, 8, 10, 12]           # coluna inicial de cada cartão (B, D, F, H, J, L)
N_PAINEL_ESTOQUE = 10                     # itens exibidos na fila de atenção do DASHBOARD

MESES_PT = '"Jan","Fev","Mar","Abr","Mai","Jun","Jul","Ago","Set","Out","Nov","Dez"'


def L(c):
    return get_column_letter(c)


# ------------------------------------------------------------------ utilidades de célula
def put(ws, ref, valor=None, fmt=None, font=None, fundo=None, ali=None, borda=None):
    c = ws[ref]
    if valor is not None:
        c.value = valor
    if fmt:
        c.number_format = fmt
    if font:
        c.font = font
    if fundo:
        c.fill = fill(fundo)
    if ali:
        c.alignment = ali
    if borda:
        c.border = borda
    return c


def ali(h="left", v="center", wrap=False, ind=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=ind)


def area(ws, r1, c1, r2, c2, fundo=None, borda_ext=None, fonte=None):
    """Preenche um retângulo e, se pedido, desenha borda externa (cor `borda_ext`)."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cel = ws.cell(r, c)
            if fundo:
                cel.fill = fill(fundo)
            if fonte:
                cel.font = fonte
            if borda_ext:
                cel.border = Border(left=lado(borda_ext) if c == c1 else None, right=lado(borda_ext) if c == c2 else None,
                                    top=lado(borda_ext) if r == r1 else None, bottom=lado(borda_ext) if r == r2 else None)


def bordas(ws, r1, c1, r2, c2, esq=None, dir_=None, topo=None, base=None):
    """Aplica lados de borda nas células de borda do retângulo (preserva os demais lados). Usar DEPOIS de mesclar:
    o openpyxl recria as células não-âncora de uma mesclagem sem estilo."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cel = ws.cell(r, c)
            b = cel.border
            cel.border = Border(left=esq if (esq and c == c1) else b.left, right=dir_ if (dir_ and c == c2) else b.right,
                                top=topo if (topo and r == r1) else b.top, bottom=base if (base and r == r2) else b.bottom)


def cartao(ws, r0, c1, rot, fml, fmt, cor, leit, sz_val, sz_rot, sz_leit, wrap_rot=True):
    """Cartão de 3 linhas (rótulo / valor / leitura) em 2 colunas, com filete de cor no topo e 'vão' entre cartões."""
    c2 = c1 + 1
    mescla(ws, r0, c1, c2, rot, font=f(sz_rot, True, CINZA_T), ali=ali("left", "center", wrap_rot, 1))
    mescla(ws, r0 + 1, c1, c2, fml, fmt=fmt, font=f(sz_val, True, NAVY), ali=ali("left", "center", False, 1))
    mescla(ws, r0 + 2, c1, c2, leit, font=f(sz_leit, False, CINZA_T), ali=ali("left", "top", True, 1))
    area(ws, r0, c1, r0 + 2, c2, fundo=BRANCO)
    bordas(ws, r0, c1, r0 + 2, c2, esq=lado(BG, "thick"), dir_=lado(BG, "thick"))
    bordas(ws, r0, c1, r0, c2, topo=lado(cor, "thick"))


def mescla(ws, r, c1, c2, valor=None, **kw):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    return put(ws, f"{L(c1)}{r}", valor, **kw)


def cab(ws, r, textos, c0=1, cor=NAVY):
    for j, t in enumerate(textos, start=c0):
        c = ws.cell(r, j, t)
        c.font, c.fill = f(10, True, "FFFFFF"), fill(cor)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(*(lado("BFBFBF"),) * 4)
    ws.row_dimensions[r].height = 30


def cel_dado(ws, r, c, valor, fmt=None, h="left", negrito=False):
    x = ws.cell(r, c, valor)
    x.font = f(10, negrito)
    x.alignment = Alignment(horizontal=h, vertical="center")
    x.border = Border(*(lado("BFBFBF"),) * 4)
    if fmt:
        x.number_format = fmt
    return x


# ================================================================== aba Analises (camada de cálculo)
def construir_analises(wb, ind, n_mat):
    """Escreve a aba Analises. Retorna o dicionário de posições usado por DASHBOARD, gráficos e testes."""
    ws = wb.create_sheet(S_ANA)
    kpi = ind["kpi_row"]
    K_UN, K_CUSTO, K_MOV, K_MATD = (f"Indicadores!$B${kpi + i}" for i in range(4))
    tot = ind["tot"]
    P = {}                                   # posições para consumo externo

    def primeira(chave, n):
        return tot[chave] - n

    ws["A1"] = "Analises – camada de cálculo do DASHBOARD (V0.2)"
    ws["A1"].font = f(14, True, NAVY)
    ws["A2"] = AVISO
    ws["A2"].font = f(9, i=True, cor="595959")
    ws["A3"] = ("Todas as células desta aba são fórmulas (nada digitado). Cadeia de rastreabilidade: tbConsumo/tbEstoque/tbMateriais -> "
                "Indicadores -> Analises -> DASHBOARD. Sem previsão: apenas soma, contagem, ordenação e percentuais do período observado.")
    ws["A3"].font = f(9, i=True, cor="595959")
    r = 6

    # ---------------------------------------------------------------- 1. mensal
    ws.cell(r, 1, "1. Evolução mensal (lê Indicadores, seção 2)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Mês (chave)", "Rótulo", "Movimentações", "Unidades consumidas", "Custo estimado",
                    "Var. unidades vs mês anterior", "Var. custo vs mês anterior"])
    i0 = primeira("mes", 3)
    m0 = r + 2
    for k in range(3):
        rr = m0 + k
        cel_dado(ws, rr, 1, f"=Indicadores!A{i0 + k}", h="center")
        cel_dado(ws, rr, 2, f"=CHOOSE(VALUE(RIGHT(A{rr},2)),{MESES_PT})&\"/\"&LEFT(A{rr},4)", h="center", negrito=True)
        cel_dado(ws, rr, 3, f"=Indicadores!B{i0 + k}", "#,##0", "right")
        cel_dado(ws, rr, 4, f"=Indicadores!C{i0 + k}", "#,##0", "right")
        cel_dado(ws, rr, 5, f"=Indicadores!D{i0 + k}", "#,##0.00", "right")
        for col, base in ((6, "D"), (7, "E")):
            if k == 0:
                cel_dado(ws, rr, col, "—", h="right")
            else:
                cel_dado(ws, rr, col, f'=IF({base}{rr - 1}=0,"n/d",{base}{rr}/{base}{rr - 1}-1)', "+0.0%;-0.0%;0.0%", "right")
    rt = m0 + 3
    cel_dado(ws, rt, 1, "Total", negrito=True)
    ws.cell(rt, 2).border = Border(*(lado("BFBFBF"),) * 4)
    for col, fmt in ((3, "#,##0"), (4, "#,##0"), (5, "#,##0.00")):
        cel_dado(ws, rt, col, f"=SUM({L(col)}{m0}:{L(col)}{rt - 1})", fmt, "right", True)
    P["mes"] = {"r0": m0, "r1": m0 + 2, "tot": rt, "hdr": r + 1}
    r = rt + 3

    # ---------------------------------------------------------------- 2/3. categoria e centro
    def bloco_dim(r, titulo, chave, n, rotulo, plural, ord_txt, kp_un, kp_custo):
        ws.cell(r, 1, titulo).font = f(11, True, NAVY)
        cab(ws, r + 1, [rotulo, "Movimentações", "Unidades consumidas", "Custo estimado", "% das unidades", "% do custo",
                        "Posição (unidades)", "Posição (custo)"])
        b0, b1 = r + 2, r + 1 + n
        ii = primeira(chave, n)
        for k in range(n):
            rr = b0 + k
            cel_dado(ws, rr, 1, f"=Indicadores!A{ii + k}", h="left")
            cel_dado(ws, rr, 2, f"=Indicadores!B{ii + k}", "#,##0", "right")
            cel_dado(ws, rr, 3, f"=Indicadores!C{ii + k}", "#,##0", "right")
            cel_dado(ws, rr, 4, f"=Indicadores!D{ii + k}", "#,##0.00", "right")
            cel_dado(ws, rr, 5, f"=IF({kp_un}=0,0,C{rr}/{kp_un})", "0.0%", "right")
            cel_dado(ws, rr, 6, f"=IF({kp_custo}=0,0,D{rr}/{kp_custo})", "0.0%", "right")
            for col, base in ((7, "C"), (8, "D")):      # posição única: RANK + desempate pela ordem original
                cel_dado(ws, rr, col, f"=RANK({base}{rr},{base}${b0}:{base}${b1},0)+COUNTIF({base}${b0}:{base}{rr},{base}{rr})-1", "0", "center")
        rt = b1 + 1
        cel_dado(ws, rt, 1, "Total", negrito=True)
        for col, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "#,##0.00"), (5, "0.0%"), (6, "0.0%")):
            cel_dado(ws, rt, col, f"=SUM({L(col)}{b0}:{L(col)}{b1})", fmt, "right", True)
        # tabelas ordenadas (lado a lado)
        s_t = rt + 2
        ws.cell(s_t, 1, f"{plural} {ord_txt} por unidades").font = f(10, True, GRAFITE)
        ws.cell(s_t, 6, f"{plural} {ord_txt} por custo").font = f(10, True, GRAFITE)
        cab(ws, s_t + 1, ["Posição", rotulo, "Unidades consumidas", "% das unidades"], 1)
        cab(ws, s_t + 1, ["Posição", rotulo, "Custo estimado", "% do custo"], 6)
        s0, s1 = s_t + 2, s_t + 1 + n
        for k in range(n):
            rr = s0 + k
            cel_dado(ws, rr, 1, f"=ROWS(A${s0}:A{rr})", "0", "center")
            cel_dado(ws, rr, 2, f"=INDEX($A${b0}:$A${b1},MATCH($A{rr},$G${b0}:$G${b1},0))")
            cel_dado(ws, rr, 3, f"=INDEX($C${b0}:$C${b1},MATCH($A{rr},$G${b0}:$G${b1},0))", "#,##0", "right")
            cel_dado(ws, rr, 4, f"=INDEX($E${b0}:$E${b1},MATCH($A{rr},$G${b0}:$G${b1},0))", "0.0%", "right")
            cel_dado(ws, rr, 6, f"=ROWS(F${s0}:F{rr})", "0", "center")
            cel_dado(ws, rr, 7, f"=INDEX($A${b0}:$A${b1},MATCH($F{rr},$H${b0}:$H${b1},0))")
            cel_dado(ws, rr, 8, f"=INDEX($D${b0}:$D${b1},MATCH($F{rr},$H${b0}:$H${b1},0))", "#,##0.00", "right")
            cel_dado(ws, rr, 9, f"=INDEX($F${b0}:$F${b1},MATCH($F{rr},$H${b0}:$H${b1},0))", "0.0%", "right")
        return {"b0": b0, "b1": b1, "tot": rt, "s0": s0, "s1": s1, "hdr_s": s_t + 1}, s1 + 3

    P["cat"], r = bloco_dim(r, "2. Categorias (lê Indicadores, seção 3)", "cat", 6, "Categoria", "Categorias", "ordenadas", K_UN, K_CUSTO)
    P["cen"], r = bloco_dim(r, "3. Centros de trabalho (lê Indicadores, seção 4)", "cen", 5, "Centro_Trabalho", "Centros", "ordenados", K_UN, K_CUSTO)

    # ---------------------------------------------------------------- 4. materiais - base
    ws.cell(r, 1, "4. Materiais – base de cálculo (uma linha por material do catálogo; lê Indicadores, seção 5)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Material_ID", "Material", "Categoria", "Criticidade", "Movimentações", "Unidades consumidas",
                    "Custo estimado", "Posição (unidades)", "Posição (custo)"])
    b0, b1 = r + 2, r + 1 + n_mat
    im = primeira("mat", n_mat)
    for k in range(n_mat):
        rr = b0 + k
        for col, src, fmt, h in ((1, "A", None, "left"), (2, "F", None, "left"), (3, "G", None, "left"), (4, "H", None, "center"),
                                 (5, "B", "#,##0", "right"), (6, "C", "#,##0", "right"), (7, "D", "#,##0.00", "right")):
            cel_dado(ws, rr, col, f"=Indicadores!{src}{im + k}", fmt, h)
        for col, base in ((8, "F"), (9, "G")):
            cel_dado(ws, rr, col, f"=RANK({base}{rr},{base}${b0}:{base}${b1},0)+COUNTIF({base}${b0}:{base}{rr},{base}{rr})-1", "0", "center")
    rt = b1 + 1
    cel_dado(ws, rt, 1, "Total", negrito=True)
    for col, fmt in ((5, "#,##0"), (6, "#,##0"), (7, "#,##0.00")):
        cel_dado(ws, rt, col, f"=SUM({L(col)}{b0}:{L(col)}{b1})", fmt, "right", True)
    P["mat"] = {"b0": b0, "b1": b1, "tot": rt}
    r = rt + 3

    # ---------------------------------------------------------------- 5. Pareto (custo e consumo)
    def pareto(r, titulo, rot_val, pct_rot, col_base, rank_col, kp, fmt, nome_serie_nucleo, nome_dem):
        ws.cell(r, 1, titulo).font = f(11, True, NAVY)
        hdr = r + 1
        cab(ws, hdr, ["Posição", "Material_ID", "Material", rot_val, pct_rot,
                      "% acumulado", "Limite (parâmetro)", "Faixa", "", ""])
        ws.cell(hdr, 9).value = nome_serie_nucleo        # cabeçalhos de série (fórmulas) usados pelas legendas dos gráficos
        ws.cell(hdr, 10).value = nome_dem
        ws.cell(hdr, 6).value = "% acumulado"
        ws.cell(hdr, 7).value = '="Limite "&ROUND(pLimitePareto*100,0)&"%"'
        p0, p1 = r + 2, r + 1 + n_mat
        for k in range(n_mat):
            rr = p0 + k
            cel_dado(ws, rr, 1, f"=ROWS(A${p0}:A{rr})", "0", "center")
            cel_dado(ws, rr, 2, f"=INDEX($A${b0}:$A${b1},MATCH($A{rr},${rank_col}${b0}:${rank_col}${b1},0))")
            cel_dado(ws, rr, 3, f"=INDEX($B${b0}:$B${b1},MATCH($A{rr},${rank_col}${b0}:${rank_col}${b1},0))")
            cel_dado(ws, rr, 4, f"=INDEX(${col_base}${b0}:${col_base}${b1},MATCH($A{rr},${rank_col}${b0}:${rank_col}${b1},0))", fmt, "right")
            cel_dado(ws, rr, 5, f"=IF({kp}=0,0,D{rr}/{kp})", "0.0%", "right")
            cel_dado(ws, rr, 6, f"=SUM(E${p0}:E{rr})", "0.0%", "right")
            cel_dado(ws, rr, 7, "=pLimitePareto", "0%", "right")
            cel_dado(ws, rr, 8, f'=IF(F{rr}-E{rr}<G{rr},"Núcleo","Cauda")', h="center")
            cel_dado(ws, rr, 9, f'=IF(H{rr}="Núcleo",D{rr},0)', fmt, "right")
            cel_dado(ws, rr, 10, f'=IF(H{rr}="Cauda",D{rr},0)', fmt, "right")
        rt = p1 + 1
        cel_dado(ws, rt, 1, "Total", negrito=True)
        cel_dado(ws, rt, 4, f"=SUM(D{p0}:D{p1})", fmt, "right", True)
        cel_dado(ws, rt, 5, f"=SUM(E{p0}:E{p1})", "0.0%", "right", True)
        return {"hdr": hdr, "p0": p0, "p1": p1, "tot": rt}, rt + 3

    P["par_c"], r = pareto(r, "5a. Pareto de custo por material", "Custo estimado", "% do custo", "G", "I", K_CUSTO, "#,##0.00",
                           '="Núcleo (até "&ROUND(pLimitePareto*100,0)&"% do custo)"', "Demais materiais")
    P["par_u"], r = pareto(r, "5b. Pareto de consumo por material", "Unidades consumidas", "% das unidades", "F", "H", K_UN, "#,##0",
                           '="Núcleo (até "&ROUND(pLimitePareto*100,0)&"% das unidades)"', "Demais materiais")

    # ---------------------------------------------------------------- 6. concentração
    ws.cell(r, 1, "6. Concentração (números do Pareto; sem juízo de valor)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Indicador", "Valor", "Detalhe", "Base de cálculo"])
    pc, pu, cat, cen = P["par_c"], P["par_u"], P["cat"], P["cen"]
    linhas = [
        ("mat80_c", '="Materiais que somam "&ROUND(pLimitePareto*100,0)&"% do custo"',
         f'=MIN(COUNTIF(D{pc["p0"]}:D{pc["p1"]},">0"),COUNTIF(F{pc["p0"]}:F{pc["p1"]},"<"&pLimitePareto)+1)', "0",
         "Menor nº de materiais cujo % acumulado de custo atinge o limite (Pareto 5a)."),
        ("pct80_c", "… em % dos materiais movimentados", f"=IF({K_MATD}=0,0,B{{mat80_c}}/{K_MATD})", "0.0%", "Linha anterior ÷ KPI 'Materiais distintos movimentados'."),
        ("topa_c", '="Custo: participação dos "&pTopA&" maiores materiais"', f"=INDEX(F{pc['p0']}:F{pc['p1']},pTopA)", "0.0%", "% acumulado do custo na posição pTopA (aba Parametros)."),
        ("topb_c", '="Custo: participação dos "&pTopB&" maiores materiais"', f"=INDEX(F{pc['p0']}:F{pc['p1']},pTopB)", "0.0%", "% acumulado do custo na posição pTopB (aba Parametros)."),
        ("mat80_u", '="Materiais que somam "&ROUND(pLimitePareto*100,0)&"% das unidades"',
         f'=MIN(COUNTIF(D{pu["p0"]}:D{pu["p1"]},">0"),COUNTIF(F{pu["p0"]}:F{pu["p1"]},"<"&pLimitePareto)+1)', "0",
         "Menor nº de materiais cujo % acumulado de unidades atinge o limite (Pareto 5b)."),
        ("pct80_u", "… em % dos materiais movimentados", f"=IF({K_MATD}=0,0,B{{mat80_u}}/{K_MATD})", "0.0%", "Linha anterior ÷ KPI 'Materiais distintos movimentados'."),
        ("topa_u", '="Unidades: participação dos "&pTopA&" maiores materiais"', f"=INDEX(F{pu['p0']}:F{pu['p1']},pTopA)", "0.0%", "% acumulado das unidades na posição pTopA."),
        ("topb_u", '="Unidades: participação dos "&pTopB&" maiores materiais"', f"=INDEX(F{pu['p0']}:F{pu['p1']},pTopB)", "0.0%", "% acumulado das unidades na posição pTopB."),
        ("cat_c", "Maior categoria em custo", f"=I{cat['s0']}", "0.0%", "Primeira linha da tabela de categorias ordenada por custo (detalhe = nome)."),
        ("cat_u", "Maior categoria em unidades", f"=D{cat['s0']}", "0.0%", "Primeira linha da tabela de categorias ordenada por unidades."),
        ("cen_c", "Maior centro de trabalho em custo", f"=I{cen['s0']}", "0.0%", "Primeira linha da tabela de centros ordenada por custo."),
        ("cen_u", "Maior centro de trabalho em unidades", f"=D{cen['s0']}", "0.0%", "Primeira linha da tabela de centros ordenada por unidades."),
    ]
    detalhes = {"cat_c": f"=G{cat['s0']}", "cat_u": f"=B{cat['s0']}", "cen_c": f"=G{cen['s0']}", "cen_u": f"=B{cen['s0']}",
                "topa_c": f"=C{pc['p0']}", "topa_u": f"=C{pu['p0']}"}
    P["conc"] = {}
    for k, (chave, _, _, _, _) in enumerate(linhas):
        P["conc"][chave] = r + 2 + k
    for k, (chave, rot, fml, fmt, base) in enumerate(linhas):
        rr = r + 2 + k
        cel_dado(ws, rr, 1, rot)
        cel_dado(ws, rr, 2, fml.format(**P["conc"]) if "{" in fml else fml, fmt, "right", True)
        cel_dado(ws, rr, 3, detalhes.get(chave))
        cel_dado(ws, rr, 4, base)
    r = r + 2 + len(linhas) + 2

    # ---------------------------------------------------------------- 7. estoque e criticidade
    ws.cell(r, 1, "7a. Estoque – base de prioridade (lê tbEstoque; Status_Estoque e Critico_Em_Risco vêm das colunas calculadas da V0.1.1)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["#", "Material_ID", "Material", "Categoria", "Criticidade", "Estoque_Atual", "Estoque_Minimo", "Status_Estoque",
                    "Critico_Em_Risco", "Atual / Mínimo", "Chave de prioridade", "Posição na fila"])
    e0, e1 = r + 2, r + 1 + n_mat
    for k in range(n_mat):
        rr = e0 + k
        cel_dado(ws, rr, 1, f"=ROWS(A${e0}:A{rr})", "0", "center")
        for col, campo, h in ((2, "Material_ID", "left"), (3, "Material", "left"), (4, "Categoria", "left"), (5, "Criticidade", "center"),
                              (6, "Estoque_Atual", "right"), (7, "Estoque_Minimo", "right"), (8, "Status_Estoque", "center"),
                              (9, "Critico_Em_Risco", "center")):
            cel_dado(ws, rr, col, f"=INDEX(tbEstoque[{campo}],$A{rr})", "#,##0" if col in (6, 7) else None, h)
        cel_dado(ws, rr, 10, f'=IF(G{rr}=0,"",F{rr}/G{rr})', "0.00", "right")
        cel_dado(ws, rr, 11, (f'=IF(H{rr}="REPOR",1,IF(H{rr}="ATENÇÃO",2,3))*1000000+IF(E{rr}="Alta",1,IF(E{rr}="Média",2,3))*10000'
                              f'+MIN(9999,ROUND(IF(G{rr}=0,9.999,F{rr}/G{rr})*1000,0))'), "0", "right")
        cel_dado(ws, rr, 12, f"=RANK(K{rr},K${e0}:K${e1},1)+COUNTIF(K${e0}:K{rr},K{rr})-1", "0", "center")
    P["est_b"] = {"e0": e0, "e1": e1}
    r = e1 + 3

    ws.cell(r, 1, "7b. Estoque – fila de prioridade (REPOR primeiro; depois ATENÇÃO; dentro do status: criticidade Alta > Média > Baixa; depois menor cobertura Atual/Mínimo)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Posição", "Material_ID", "Material", "Categoria", "Criticidade", "Estoque_Atual", "Estoque_Minimo",
                    "Atual - Mínimo", "Atual / Mínimo", "Status_Estoque", "Critico_Em_Risco"])
    q0, q1 = r + 2, r + 1 + n_mat
    for k in range(n_mat):
        rr = q0 + k
        cel_dado(ws, rr, 1, f"=ROWS(A${q0}:A{rr})", "0", "center")
        for col, src, fmt, h in ((2, "B", None, "left"), (3, "C", None, "left"), (4, "D", None, "left"), (5, "E", None, "center"),
                                 (6, "F", "#,##0", "right"), (7, "G", "#,##0", "right"), (9, "J", "0.00", "right"),
                                 (10, "H", None, "center"), (11, "I", None, "center")):
            cel_dado(ws, rr, col, f"=INDEX(${src}${e0}:${src}${e1},MATCH($A{rr},$L${e0}:$L${e1},0))", fmt, h)
        cel_dado(ws, rr, 8, f"=F{rr}-G{rr}", "+#,##0;-#,##0;0", "right")
    P["est_q"] = {"q0": q0, "q1": q1}
    ws.conditional_formatting.add(f"J{q0}:J{q1}", CellIsRule(operator="equal", formula=['"REPOR"'], fill=FILL_REPOR, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"J{q0}:J{q1}", CellIsRule(operator="equal", formula=['"ATENÇÃO"'], fill=FILL_ATEN, font=Font(name=FONTE, bold=True, color="7F6000")))
    ws.conditional_formatting.add(f"H{e0}:H{e1}", CellIsRule(operator="equal", formula=['"REPOR"'], fill=FILL_REPOR, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"H{e0}:H{e1}", CellIsRule(operator="equal", formula=['"ATENÇÃO"'], fill=FILL_ATEN, font=Font(name=FONTE, bold=True, color="7F6000")))
    r = q1 + 3

    ws.cell(r, 1, "7c. Matriz criticidade × status (contagem de itens; lê tbEstoque)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Criticidade", "REPOR", "ATENÇÃO", "OK", "Total"])
    x0 = r + 2
    for k, ref_c in enumerate(("Parametros!$C$13", "Parametros!$C$12", "Parametros!$C$11")):     # Alta, Média, Baixa
        rr = x0 + k
        cel_dado(ws, rr, 1, f"={ref_c}", negrito=True)
        for col in (2, 3, 4):
            cel_dado(ws, rr, col, f"=COUNTIFS(tbEstoque[Criticidade],$A{rr},tbEstoque[Status_Estoque],{L(col)}${r + 1})", "#,##0", "center")
        cel_dado(ws, rr, 5, f"=SUM(B{rr}:D{rr})", "#,##0", "center", True)
    rt = x0 + 3
    cel_dado(ws, rt, 1, "Total", negrito=True)
    for col in (2, 3, 4, 5):
        cel_dado(ws, rt, col, f"=SUM({L(col)}{x0}:{L(col)}{rt - 1})", "#,##0", "center", True)
    P["mtx"] = {"hdr": r + 1, "x0": x0, "tot": rt}
    r = rt + 3

    ws.cell(r, 1, "7d. Painel de atenção do DASHBOARD (primeiros itens da fila que NÃO estão OK)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Posição", "Material", "Criticidade", "Estoque atual", "Estoque mínimo", "Status_Estoque", "Critico_Em_Risco"])
    n0, n1 = r + 2, r + 1 + N_PAINEL_ESTOQUE
    for k in range(N_PAINEL_ESTOQUE):
        rr = n0 + k
        cel_dado(ws, rr, 1, f"=ROWS(A${n0}:A{rr})", "0", "center")
        for col, src, fmt, h in ((2, "C", None, "left"), (3, "E", None, "center"), (4, "F", "#,##0", "right"),
                                 (5, "G", "#,##0", "right"), (6, "J", None, "center"), (7, "K", None, "center")):
            cel_dado(ws, rr, col, f'=IF($J{q0 + k}="OK","",{src}{q0 + k})', fmt, h)
    P["pnl"] = {"n0": n0, "n1": n1}
    rr = n1 + 1
    cel_dado(ws, rr, 1, "Itens fora do OK", negrito=True)
    cel_dado(ws, rr, 2, f"=B{P['mtx']['tot']}+C{P['mtx']['tot']}", "#,##0", "right", True)
    cel_dado(ws, rr + 1, 1, "Não exibidos no painel", negrito=True)
    cel_dado(ws, rr + 1, 2, f"=MAX(0,B{rr}-{N_PAINEL_ESTOQUE})", "#,##0", "right", True)
    P["pnl"]["fora"], P["pnl"]["oculto"] = rr, rr + 1
    r = rr + 4

    P["conf_ini"] = r
    return ws, P, r


def conferencias(ws, P, ind, r, n_mat):
    """Bloco 8: conferências da camada analítica (cada linha compara dois números que devem coincidir)."""
    kpi = ind["kpi_row"]
    KU, KC, KM, KD, KR, KK = (f"Indicadores!$B${kpi + i}" for i in range(6))
    m, cat, cen, pc, pu, mt, pn = P["mes"], P["cat"], P["cen"], P["par_c"], P["par_u"], P["mtx"], P["pnl"]
    ws.cell(r, 1, "8. Conferências da camada analítica (Valor A deve ser igual ao Valor B)").font = f(11, True, NAVY)
    cab(ws, r + 1, ["Conferência", "Valor A", "Valor B", "Diferença", "Status"])
    dash_vals = [f"{S_DASH}!{L(c)}{KPI_VALOR}" for c in KPI_COLS]
    itens = [
        ("Σ unidades por mês = KPI Unidades consumidas", f"=D{m['tot']}", f"={KU}", "#,##0.00"),
        ("Σ custo por mês = KPI Custo estimado", f"=E{m['tot']}", f"={KC}", "#,##0.00"),
        ("Σ movimentações por mês = KPI Movimentações", f"=C{m['tot']}", f"={KM}", "#,##0.00"),
        ("Σ unidades por categoria (ordenada) = KPI", f"=SUM(C{cat['s0']}:C{cat['s1']})", f"={KU}", "#,##0.00"),
        ("Σ custo por categoria (ordenada) = KPI", f"=SUM(H{cat['s0']}:H{cat['s1']})", f"={KC}", "#,##0.00"),
        ("Σ unidades por centro (ordenada) = KPI", f"=SUM(C{cen['s0']}:C{cen['s1']})", f"={KU}", "#,##0.00"),
        ("Σ custo por centro (ordenada) = KPI", f"=SUM(H{cen['s0']}:H{cen['s1']})", f"={KC}", "#,##0.00"),
        ("Σ custo do Pareto de custo = KPI Custo estimado", f"=D{pc['tot']}", f"={KC}", "#,##0.00"),
        ("Σ unidades do Pareto de consumo = KPI Unidades", f"=D{pu['tot']}", f"={KU}", "#,##0.00"),
        ("% acumulado final do Pareto de custo = 100%", f"=F{pc['p1']}", "=1", "0.0000"),
        ("% acumulado final do Pareto de consumo = 100%", f"=F{pu['p1']}", "=1", "0.0000"),
        ("Posições únicas nos rankings de materiais (unidades e custo) e da fila de estoque (esperado: 3 × nº de materiais)",
         (f"=SUMPRODUCT(1/COUNTIF(H{P['mat']['b0']}:H{P['mat']['b1']},H{P['mat']['b0']}:H{P['mat']['b1']}))"
          f"+SUMPRODUCT(1/COUNTIF(I{P['mat']['b0']}:I{P['mat']['b1']},I{P['mat']['b0']}:I{P['mat']['b1']}))"
          f"+SUMPRODUCT(1/COUNTIF(L{P['est_b']['e0']}:L{P['est_b']['e1']},L{P['est_b']['e0']}:L{P['est_b']['e1']}))"),
         f"=3*ROWS(tbMateriais[Material_ID])", "#,##0.00"),
        ("Matriz criticidade × status: total = linhas de tbEstoque", f"=E{mt['tot']}", "=ROWS(tbEstoque[Material_ID])", "#,##0.00"),
        ("Matriz: coluna REPOR = KPI Itens abaixo ou iguais ao mínimo", f"=B{mt['tot']}", f"={KR}", "#,##0.00"),
        ("Matriz: Alta × REPOR = KPI Itens críticos", f"=B{mt['x0']}", f"={KK}", "#,##0.00"),
        ("Fila de prioridade: itens REPOR listados = KPI", f'=COUNTIF(J{P["est_q"]["q0"]}:J{P["est_q"]["q1"]},"REPOR")', f"={KR}", "#,##0.00"),
        ("Fila de prioridade: nenhum item REPOR depois de um item ATENÇÃO/OK (nº de inversões)",
         f'=SUMPRODUCT((J{P["est_q"]["q0"]}:J{P["est_q"]["q1"] - 1}<>"REPOR")*(J{P["est_q"]["q0"] + 1}:J{P["est_q"]["q1"]}="REPOR"))', "=0", "#,##0.00"),
        ("Cartões do DASHBOARD = 6 KPIs de Indicadores (Σ |diferenças|)",
         "=" + "+".join(f"ABS({d}-{k})" for d, k in zip(dash_vals, (KU, KC, KM, KD, KR, KK))), "=0", "#,##0.00"),
    ]
    c0 = r + 2
    for k, (rot, fa, fb, fmt) in enumerate(itens):
        rr = c0 + k
        cel_dado(ws, rr, 1, rot)
        cel_dado(ws, rr, 2, fa, fmt, "right")
        cel_dado(ws, rr, 3, fb, fmt, "right")
        cel_dado(ws, rr, 4, f"=B{rr}-C{rr}", fmt, "right")
        cel_dado(ws, rr, 5, f'=IF(ABS(D{rr})<0.005,"OK","DIVERGE")', h="center", negrito=True)
    c1 = c0 + len(itens) - 1
    ws.conditional_formatting.add(f"E{c0}:E{c1}", CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_OK, font=Font(name=FONTE, bold=True, color="1E6B3A")))
    ws.conditional_formatting.add(f"E{c0}:E{c1}", CellIsRule(operator="equal", formula=['"DIVERGE"'], fill=FILL_REPOR, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.cell(c1 + 2, 1, "Conferências OK").font = f(10, True)
    rs = c1 + 2
    cel_dado(ws, rs, 2, f'=COUNTIF(E{c0}:E{c1},"OK")&" de "&ROWS(E{c0}:E{c1})', h="right", negrito=True)
    P["conf"] = {"c0": c0, "c1": c1, "resumo": rs}

    # aparência da aba
    for j, w in enumerate([34, 40, 30, 22, 20, 20, 20, 18, 18, 22, 20, 16], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.column_dimensions["A"].width = 46
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.sheet_properties.tabColor = CINZA_T
    return P


# ================================================================== gráficos
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def _txt(sz=800, cor="44546A", bold=False, rot=None):
    cp = CharacterProperties(latin=DFont(typeface=FONTE), sz=sz, b=bold, solidFill=cor)
    corpo = RichTextProperties(rot=rot, vert="horz") if rot is not None else RichTextProperties()
    return RichText(bodyPr=corpo, p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp, r=[])])


def GraficoBarra():
    """BarChart que declara `autoTitleDeleted` (sem isso o Excel cria título automático em gráficos de 1 série).

    Não usa subclasse: subclasses de Serialisable do openpyxl não herdam __elements__ e serializam vazias.
    O ajuste é feito na instância, interceptando _write().
    """
    g = BarChart()
    original = g._write

    def _write():
        arvore = original()
        ch = arvore.find("chart")                     # o openpyxl serializa sem prefixo (namespace no atributo xmlns)
        if ch is None:
            ch = arvore.find(_C + "chart")
        if ch is None:
            raise RuntimeError("estrutura de gráfico inesperada: elemento 'chart' não encontrado")
        if ch.find("autoTitleDeleted") is None and ch.find(_C + "autoTitleDeleted") is None:
            ch.insert(0, Element("autoTitleDeleted", val="1"))
        for el in arvore.iter():                     # formatos numéricos dos rótulos: declarar sourceLinked="0" explicitamente
            if el.tag.split("}")[-1] == "numFmt" and "sourceLinked" not in el.attrib:
                el.set("sourceLinked", "0")
        return arvore

    g._write = _write
    return g


def _base_grafico(g):
    g.roundedCorners = False
    gp = GraphicalProperties(solidFill=BRANCO)
    gp.line = LineProperties(noFill=True)
    g.graphical_properties = gp
    g.title = None
    g.style = 2


def _serie_barra(s, cor):
    s.graphicalProperties = GraphicalProperties(solidFill=cor)
    s.graphicalProperties.line = LineProperties(solidFill=cor)
    s.invertIfNegative = False


def _rotulos(s, fmt, pos="outEnd", sz=800, cor="37474F"):
    d = DataLabelList()
    d.showVal, d.showSerName, d.showCatName, d.showLegendKey, d.showPercent = True, False, False, False, False
    d.numFmt, d.position, d.txPr = fmt, pos, _txt(sz, cor)
    s.dLbls = d


def _eixos(g, valor_visivel=False, fmt="#,##0", grade=True):
    g.x_axis.delete = False
    g.x_axis.txPr = _txt(800)
    g.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINHA))
    g.x_axis.majorTickMark = "none"
    g.y_axis.delete = not valor_visivel
    g.y_axis.number_format = fmt
    g.y_axis.txPr = _txt(800, CINZA_T)
    g.y_axis.majorTickMark = "none"
    g.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    if grade and valor_visivel:
        g.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E8EC"))
    else:
        g.y_axis.majorGridlines = None


def grafico_barras(ws_dados, cat_col, val_col, r0, r1, cor, fmt, horizontal, hdr_row=None):
    """Barras simples (uma série) com rótulos de dados; horizontal = ranking (maior no topo)."""
    g = GraficoBarra()
    g.type = "bar" if horizontal else "col"
    g.grouping = "clustered"
    g.gapWidth = 45 if horizontal else 70
    g.legend = None
    _base_grafico(g)
    g.add_data(Reference(ws_dados, min_col=val_col, min_row=hdr_row or r0, max_row=r1), titles_from_data=bool(hdr_row))
    g.set_categories(Reference(ws_dados, min_col=cat_col, min_row=r0, max_row=r1))
    _serie_barra(g.series[0], cor)
    _rotulos(g.series[0], fmt)
    _eixos(g, valor_visivel=False)
    if horizontal:
        g.x_axis.scaling.orientation = "maxMin"
        # margem à direita para o rótulo do maior valor não encostar na borda (o eixo de valores está oculto)
        g.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge", x=0.2, y=0.03, w=0.62, h=0.94))
    return g


def grafico_pareto(ws_dados, pr, cor_nucleo, cor_cauda, fmt_eixo):
    """Pareto: colunas (núcleo/demais, sobrepostas) + % acumulado e limite no eixo secundário (0-100%)."""
    p0, p1, h = pr["p0"], pr["p1"], pr["hdr"]
    g = GraficoBarra()
    g.type, g.grouping, g.overlap, g.gapWidth = "col", "clustered", 100, 35
    _base_grafico(g)
    for col, cor in ((9, cor_nucleo), (10, cor_cauda)):
        g.add_data(Reference(ws_dados, min_col=col, min_row=h, max_row=p1), titles_from_data=True)
    g.set_categories(Reference(ws_dados, min_col=3, min_row=p0, max_row=p1))
    _serie_barra(g.series[0], cor_nucleo)
    _serie_barra(g.series[1], cor_cauda)
    _eixos(g, valor_visivel=True, fmt=fmt_eixo)
    g.x_axis.txPr = _txt(700, "44546A", rot=-5400000)
    g.x_axis.tickLblPos = "low"
    g.y_axis.crosses = "min"

    ln = LineChart()
    for col in (6, 7):
        ln.add_data(Reference(ws_dados, min_col=col, min_row=h, max_row=p1), titles_from_data=True)
    ln.set_categories(Reference(ws_dados, min_col=3, min_row=p0, max_row=p1))
    acum, lim = ln.series
    acum.graphicalProperties = GraphicalProperties()
    acum.graphicalProperties.line = LineProperties(solidFill=GRAFITE, w=22225)
    acum.marker.symbol, acum.marker.size = "circle", 5
    acum.marker.graphicalProperties = GraphicalProperties(solidFill=GRAFITE)
    acum.marker.graphicalProperties.line = LineProperties(solidFill=GRAFITE)
    acum.smooth = False
    lim.graphicalProperties = GraphicalProperties()
    lim.graphicalProperties.line = LineProperties(solidFill=VERM, w=15875, prstDash="dash")
    lim.marker.symbol = "none"
    lim.smooth = False
    ln.y_axis.axId = 200
    ln.y_axis.crosses = "max"
    ln.y_axis.scaling.min, ln.y_axis.scaling.max = 0, 1
    ln.y_axis.majorUnit = 0.2
    ln.y_axis.number_format = "0%"
    ln.y_axis.majorGridlines = None
    ln.y_axis.delete = False
    ln.y_axis.txPr = _txt(800, CINZA_T)
    ln.y_axis.majorTickMark = "none"
    ln.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    g += ln
    g.legend.position = "b"
    g.legend.txPr = _txt(800, "44546A")
    return g


def grafico_estoque(ws_dados, pn):
    """Estoque atual × mínimo dos itens da fila de atenção (barras horizontais agrupadas)."""
    n0, n1 = pn["n0"], pn["n1"]
    g = GraficoBarra()
    g.type, g.grouping, g.gapWidth, g.overlap = "bar", "clustered", 55, -5
    _base_grafico(g)
    for col in (4, 5):
        g.add_data(Reference(ws_dados, min_col=col, min_row=n0 - 1, max_row=n1), titles_from_data=True)
    g.set_categories(Reference(ws_dados, min_col=2, min_row=n0, max_row=n1))
    _serie_barra(g.series[0], AZUL)
    _serie_barra(g.series[1], CINZA_M)
    for s in g.series:
        _rotulos(s, "#,##0", sz=700)
    _eixos(g, valor_visivel=False)
    g.x_axis.scaling.orientation = "maxMin"
    g.x_axis.txPr = _txt(800, "222222")
    g.legend.position = "b"
    g.legend.txPr = _txt(800, "44546A")
    return g


def ancorar(ws, g, r1, c1, r2, c2):
    """Ancora o gráfico exatamente sobre o retângulo de células (r1,c1)-(r2,c2), 1-based e inclusivo."""
    g.anchor = TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=c1 - 1, row=r1 - 1), to=AnchorMarker(col=c2, row=r2))
    ws.add_chart(g)


# ================================================================== aba DASHBOARD
def construir_dashboard(wb, ws, P, ind):
    ana = wb[S_ANA]
    kpi = ind["kpi_row"]
    KU, KC, KM, KD, KR, KK = (f"Indicadores!$B${kpi + i}" for i in range(6))
    m, cat, cen, pc, pu, mt, pn, cf, co = P["mes"], P["cat"], P["cen"], P["par_c"], P["par_u"], P["mtx"], P["pnl"], P["conf"], P["conc"]
    A = S_ANA + "!"

    # ---- grade
    ws.column_dimensions["A"].width = 1.7
    ws.column_dimensions["N"].width = 1.7
    for c in range(COL_INI, COL_FIM + 1):
        ws.column_dimensions[L(c)].width = LARG_COL
    ULT = 98
    area(ws, 1, 1, ULT, 14, fundo=BG)
    alturas = {1: 5, 2: 34, 3: 18, 4: 20, 5: 8, KPI_LABEL: 27, KPI_VALOR: 38, KPI_LEIT: 28, 9: 8}
    for rr, h in alturas.items():
        ws.row_dimensions[rr].height = h

    # ---- cabeçalho
    area(ws, 1, 1, 1, 14, fundo=COBRE)
    area(ws, 2, 1, 3, 14, fundo=NAVY_D)
    mescla(ws, 2, 2, 9, "LOBO INSIGHTS INDUSTRIAL", font=f(20, True, "FFFFFF"), ali=ali("left", "center"))
    mescla(ws, 3, 2, 9, "Visão Executiva | Consumo, Custos e Estoque", font=f(11, cor="C9D6E2"), ali=ali("left", "center"))
    mescla(ws, 2, 10, 13, f'="Período: "&{A}B{m["r0"]}&" a "&{A}B{m["r1"]}', font=f(10, True, "FFFFFF"), ali=ali("right", "center"))
    mescla(ws, 3, 10, 13, "Dados 100% sintéticos · sem previsão · não é recomendação operacional", font=f(8, cor="C9D6E2", i=True), ali=ali("right", "center"))
    area(ws, 4, 1, 4, 14, fundo="E4EAF0")
    links = [("Indicadores (KPIs)", "Indicadores!A1", 2, 3), ("Analises (cálculos)", "Analises!A1", 4, 5), ("Estoque (dados)", "Estoque!A1", 6, 7),
             ("Qualidade dos dados", "Qualidade!A1", 8, 9), ("Parametros", "Parametros!A1", 10, 11), ("LEIA_ME", "LEIA_ME!A1", 12, 13)]
    for txt, loc, c1, c2 in links:
        cel = mescla(ws, 4, c1, c2, txt, font=f(9, False, "1F5FA8", u="single"), ali=ali("center", "center"))
        cel.hyperlink = Hyperlink(ref=cel.coordinate, location=loc, display=txt)

    # ---- cartões KPI
    kpis = [
        ("Unidades consumidas", f"={KU}", "#,##0", AZUL,
         f'="Média de "&ROUND({KU}/{KM},1)&" unidades por movimentação"'),
        ("Custo estimado consumido", f"={KC}", "#,##0.00", COBRE,
         f'="Média de "&FIXED({KC}/{KM},2)&" por movimentação · valor fictício"'),
        ("Movimentações", f"={KM}", "#,##0", CINZA_T,
         f'="Registros de consumo em "&ROWS(lstMeses)&" meses ("&{A}B{m["r0"]}&" a "&{A}B{m["r1"]}&")"'),
        ("Materiais distintos movimentados", f"={KD}", "#,##0", CINZA_T,
         f'="de "&ROWS(tbMateriais[Material_ID])&" materiais do catálogo ("&ROUND({KD}/ROWS(tbMateriais[Material_ID])*100,0)&"%)"'),
        ("Itens abaixo ou iguais ao mínimo", f"={KR}", "#,##0", VERM,
         f'="de "&ROWS(tbEstoque[Material_ID])&" itens · mais "&COUNTIF(tbEstoque[Status_Estoque],"ATENÇÃO")&" em ATENÇÃO"'),
        ("Itens críticos abaixo ou iguais ao mínimo", f"={KK}", "#,##0", VERM_E,
         f'="de "&COUNTIF(tbEstoque[Criticidade],"Alta")&" itens de criticidade Alta"'),
    ]
    for (rot, fml, fmt, cor, leit), c1 in zip(kpis, KPI_COLS):
        cartao(ws, KPI_LABEL, c1, rot, fml, fmt, cor, leit, 24, 8.5, 8)
    for c1 in KPI_COLS[4:]:
        ws.conditional_formatting.add(f"{L(c1)}{KPI_VALOR}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONTE, size=24, bold=True, color=VERM)))

    # ---- helpers de seção / painel
    estado = {"r": 10, "quebras": []}

    def secao(num, titulo, nota=None):
        rr = estado["r"]
        ws.row_dimensions[rr].height = 22
        mescla(ws, rr, 2, 8, f"{num} · {titulo}", font=f(11, True, NAVY), ali=ali("left", "bottom"))
        mescla(ws, rr, 9, 13, nota, font=f(8, i=True, cor=CINZA_T), ali=ali("right", "bottom"))
        for c in range(2, 14):
            ws.cell(rr, c).border = Border(bottom=lado(NAVY, "medium"))
        estado["r"] += 1
        return rr

    def painel_titulo(rr, c1, c2, titulo):
        ws.row_dimensions[rr].height = 20
        area(ws, rr, c1, rr, c2, fundo=BRANCO)
        mescla(ws, rr, c1, c2, titulo, font=f(10, True, GRAFITE), ali=ali("left", "center", False, 1))

    def moldura(r1, c1, r2, c2):
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                cel = ws.cell(rr, cc)
                b = cel.border
                cel.border = Border(left=lado(LINHA) if cc == c1 else b.left, right=lado(LINHA) if cc == c2 else b.right,
                                    top=lado(LINHA) if rr == r1 else b.top, bottom=lado(LINHA) if rr == r2 else b.bottom)

    def leitura(rr, c1, c2, fml, h=28):
        ws.row_dimensions[rr].height = h
        area(ws, rr, c1, rr, c2, fundo=BRANCO)
        mescla(ws, rr, c1, c2, fml, font=f(8.5, False, CINZA_T), ali=ali("left", "center", True, 1))

    def secao_duas(num, titulo, nota, tit_esq, tit_dir, g_esq, g_dir, leit_esq, leit_dir, linhas_graf):
        rs = secao(num, titulo, nota)
        rt = rs + 1
        painel_titulo(rt, 2, 7, tit_esq)
        painel_titulo(rt, 8, 13, tit_dir)
        g1, g2 = rt + 1, rt + linhas_graf
        area(ws, g1, 2, g2, 13, fundo=BRANCO)
        for k in range(linhas_graf):
            ws.row_dimensions[g1 + k].height = 15
        ancorar(ws, g_esq, g1, 2, g2, 7)
        ancorar(ws, g_dir, g1, 8, g2, 13)
        rl = g2 + 1
        leitura(rl, 2, 7, leit_esq)
        leitura(rl, 8, 13, leit_dir)
        moldura(rt, 2, rl, 7)
        moldura(rt, 8, rl, 13)
        ws.row_dimensions[rl + 1].height = 8
        estado["r"] = rl + 2
        return rs, rl

    # ---- 1. evolução mensal
    mm, r0m, r1m, tm = ana, m["r0"], m["r1"], m["tot"]

    def txt_var(col_var, r_a, r_b):
        return (f'IF(ISNUMBER({A}{col_var}{r_b}),IF({A}{col_var}{r_b}>=0,"+","")&ROUND({A}{col_var}{r_b}*100,1)&"%","n/d")')

    def leit_mes(col_val, col_var, un, fmt_txt):
        rmax = f'INDEX({A}$B${r0m}:$B${r1m},MATCH(MAX({A}${col_val}${r0m}:${col_val}${r1m}),{A}${col_val}${r0m}:${col_val}${r1m},0))'
        return ('="Maior "&"' + un + '"&": "&' + rmax + f'&" ("&{fmt_txt}&")"&CHAR(10)&'
                f'{A}B{r0m + 1}&" vs "&{A}B{r0m}&": "&{txt_var(col_var, r0m, r0m + 1)}&"  ·  "&{A}B{r0m + 2}&" vs "&{A}B{r0m + 1}&": "&{txt_var(col_var, r0m, r0m + 2)}')

    fmt_un = f'FIXED(MAX({A}$D${r0m}:$D${r1m}),0)&" un."'
    fmt_cu = f'FIXED(MAX({A}$E${r0m}:$E${r1m}),2)'
    g_u = grafico_barras(ana, 2, 4, r0m, r1m, AZUL, "#,##0", False)
    g_c = grafico_barras(ana, 2, 5, r0m, r1m, COBRE, "#,##0.00", False)
    rs1, rl1 = secao_duas("1", "Evolução mensal", "Variações entre meses vêm dos dados sintéticos; não indicam tendência.",
                          "Consumo mensal · unidades", "Custo estimado mensal · valor fictício", g_u, g_c,
                          leit_mes("D", "F", "consumo", fmt_un), leit_mes("E", "G", "custo", fmt_cu), 9)

    # ---- 2. categorias / 3. centros
    def sec_dim(num, titulo, d, plural_min, singular):
        g_un = grafico_barras(ana, 2, 3, d["s0"], d["s1"], AZUL, "#,##0", True)
        g_cu = grafico_barras(ana, 7, 8, d["s0"], d["s1"], COBRE, "#,##0.00", True)
        n = d["s1"] - d["s0"] + 1
        leit_u = (f'="Maior consumo: "&{A}B{d["s0"]}&" - "&FIXED({A}C{d["s0"]},0)&" un. ("&ROUND({A}D{d["s0"]}*100,1)&"% do total)"&CHAR(10)&'
                  f'"Menor: "&{A}B{d["s1"]}&" - "&FIXED({A}C{d["s1"]},0)&" un. ("&ROUND({A}D{d["s1"]}*100,1)&"%)"')
        leit_c = (f'="Maior custo: "&{A}G{d["s0"]}&" - "&FIXED({A}H{d["s0"]},2)&" ("&ROUND({A}I{d["s0"]}*100,1)&"% do total)"&CHAR(10)&'
                  f'"Menor: "&{A}G{d["s1"]}&" - "&FIXED({A}H{d["s1"]},2)&" ("&ROUND({A}I{d["s1"]}*100,1)&"%)"')
        return secao_duas(num, titulo, "Ordenado do maior para o menor; % sobre o total do período.",
                          f"{plural_min} por consumo · unidades", f"{plural_min} por custo · valor fictício", g_un, g_cu, leit_u, leit_c,
                          8 if n <= 5 else 9)

    rs2, rl2 = sec_dim("2", "Categorias", cat, "Categorias", "categoria")
    estado["quebras"].append(rl2 + 1)         # quebra de página: pág. 1 = cabeçalho, KPIs, mensal e categorias
    sec_dim("3", "Centros de trabalho", cen, "Centros de trabalho", "centro")

    # ---- 4. materiais (Pareto)
    rs4 = secao("4", "Materiais · Pareto", "Colunas = valor do material; linha = % acumulado; tracejado = limite (aba Parametros).")
    rt = rs4 + 1
    painel_titulo(rt, 2, 7, "Pareto de custo por material · valor fictício")
    painel_titulo(rt, 8, 13, "Pareto de consumo por material · unidades")
    NG = 19
    g1, g2 = rt + 1, rt + NG
    area(ws, g1, 2, g2, 13, fundo=BRANCO)
    for k in range(NG):
        ws.row_dimensions[g1 + k].height = 15
    ancorar(ws, grafico_pareto(ana, pc, COBRE, COR_COBRE_CLARO, "#,##0"), g1, 2, g2, 7)
    ancorar(ws, grafico_pareto(ana, pu, AZUL, COR_AZUL_CLARO, "#,##0"), g1, 8, g2, 13)
    moldura(rt, 2, g2, 7)
    moldura(rt, 8, g2, 13)
    ws.row_dimensions[g2 + 1].height = 8
    # cartões de concentração (6 mini-cartões, mesma grade dos KPIs)
    rc = g2 + 2
    for rr, h in ((rc, 18), (rc + 1, 32), (rc + 2, 26)):
        ws.row_dimensions[rr].height = h
    mini = [
        (f"={A}A{co['mat80_c']}", f'={A}B{co["mat80_c"]}&" de "&{KD}', None, COBRE,
         f'=ROUND({A}B{co["pct80_c"]}*100,0)&"% dos materiais movimentados"'),
        (f'="Custo · "&pTopA&" maiores materiais"', f"={A}B{co['topa_c']}", "0.0%", COBRE, f'={A}C{co["topa_c"]}&" é o maior em custo"'),
        (f'="Custo · "&pTopB&" maiores materiais"', f"={A}B{co['topb_c']}", "0.0%", COBRE, f'="do custo estimado total de "&FIXED({KC},2)'),
        (f"={A}A{co['mat80_u']}", f'={A}B{co["mat80_u"]}&" de "&{KD}', None, AZUL,
         f'=ROUND({A}B{co["pct80_u"]}*100,0)&"% dos materiais movimentados"'),
        (f'="Consumo · "&pTopA&" maiores materiais"', f"={A}B{co['topa_u']}", "0.0%", AZUL, f'={A}C{co["topa_u"]}&" é o maior em unidades"'),
        (f'="Consumo · "&pTopB&" maiores materiais"', f"={A}B{co['topb_u']}", "0.0%", AZUL, f'="das "&FIXED({KU},0)&" unidades consumidas"'),
    ]
    for (rot, fml, fmt, cor, sub), c1 in zip(mini, KPI_COLS):
        cartao(ws, rc, c1, rot, fml, fmt, cor, sub, 16, 8, 7.5)
    ws.row_dimensions[rc + 3].height = 8
    estado["quebras"].append(rc + 3)
    estado["r"] = rc + 4

    # ---- 5. estoque e criticidade
    rs5 = secao("5", "Estoque e criticidade · o que precisa de atenção", "Regra do estado atual (não é previsão); parâmetros na aba Parametros.")
    rt = rs5 + 1
    painel_titulo(rt, 2, 7, "Fila de atenção · itens em REPOR e ATENÇÃO")
    painel_titulo(rt, 8, 13, "Estoque atual × mínimo · itens da fila")
    rh = rt + 1                                     # cabeçalho da lista
    ws.row_dimensions[rh].height = 20
    area(ws, rh, 2, rh, 7, fundo="E4EAF0")
    for txt, c1, c2 in (("Material", 2, 3), ("Criticidade", 4, 4), ("Atual", 5, 5), ("Mínimo", 6, 6), ("Status", 7, 7)):
        mescla(ws, rh, c1, c2, txt, font=f(8.5, True, GRAFITE), ali=ali("left" if c1 == 2 else "center", "center", False, 1 if c1 == 2 else 0))
    l0 = rh + 1
    for k in range(N_PAINEL_ESTOQUE):
        rr = l0 + k
        ws.row_dimensions[rr].height = 17
        area(ws, rr, 2, rr, 7, fundo=BRANCO)
        for col_src, c1, c2, fmt, h in (("B", 2, 3, None, "left"), ("C", 4, 4, None, "center"), ("D", 5, 5, "#,##0", "center"),
                                        ("E", 6, 6, "#,##0", "center"), ("F", 7, 7, None, "center")):
            mescla(ws, rr, c1, c2, f"={A}{col_src}{pn['n0'] + k}", fmt=fmt, font=f(9), ali=ali(h, "center", False, 1 if c1 == 2 else 0))
        for cc in range(2, 8):
            ws.cell(rr, cc).border = Border(bottom=lado("ECEFF2"))
    l1 = l0 + N_PAINEL_ESTOQUE - 1
    ws.conditional_formatting.add(f"G{l0}:G{l1}", CellIsRule(operator="equal", formula=['"REPOR"'], fill=FILL_REPOR, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"G{l0}:G{l1}", CellIsRule(operator="equal", formula=['"ATENÇÃO"'], fill=FILL_ATEN, font=Font(name=FONTE, bold=True, color="7F6000")))
    ws.conditional_formatting.add(f"D{l0}:D{l1}", CellIsRule(operator="equal", formula=['"Alta"'], font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"B{l0}:C{l1}", FormulaRule(formula=[f'AND($D{l0}="Alta",$G{l0}="REPOR")'], font=Font(name=FONTE, bold=True, color="9C0006")))
    # matriz
    rmt = l1 + 1
    ws.row_dimensions[rmt].height = 6
    area(ws, rmt, 2, rmt, 7, fundo=BRANCO)
    rmh = rmt + 1
    ws.row_dimensions[rmh].height = 20
    area(ws, rmh, 2, rmh, 7, fundo="E4EAF0")
    mescla(ws, rmh, 2, 3, "Criticidade × status", font=f(8.5, True, GRAFITE), ali=ali("left", "center", False, 1))
    for txt_h, cc in (("REPOR", 4), ("ATENÇÃO", 5), ("OK", 6), ("Total", 7)):
        put(ws, f"{L(cc)}{rmh}", f"={A}{L(cc - 2)}{mt['hdr']}", font=f(8.5, True, GRAFITE), ali=ali("center"))
    rx0 = rmh + 1
    for k in range(4):
        rr = rx0 + k
        ws.row_dimensions[rr].height = 18
        area(ws, rr, 2, rr, 7, fundo=BRANCO)
        mescla(ws, rr, 2, 3, f"={A}A{mt['x0'] + k}" if k < 3 else "Total", font=f(9, True), ali=ali("left", "center", False, 1))
        for cc in range(4, 8):
            put(ws, f"{L(cc)}{rr}", f"={A}{L(cc - 2)}{mt['x0'] + k}", fmt="#,##0", font=f(9, k == 3 or cc == 7), ali=ali("center"))
        for cc in range(2, 8):
            ws.cell(rr, cc).border = Border(bottom=lado("ECEFF2"))
    rx1 = rx0 + 2
    ws.conditional_formatting.add(f"D{rx0}:D{rx1}", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_REPOR, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws.conditional_formatting.add(f"E{rx0}:E{rx1}", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_ATEN, font=Font(name=FONTE, bold=True, color="7F6000")))
    ws.conditional_formatting.add(f"F{rx0}:F{rx1}", CellIsRule(operator="greaterThan", formula=["0"], fill=FILL_OK, font=Font(name=FONTE, bold=True, color="1E6B3A")))
    # gráfico à direita
    g_top, g_bot = rh, rx0 + 3
    area(ws, g_top, 8, g_bot, 13, fundo=BRANCO)
    ancorar(ws, grafico_estoque(ana, pn), g_top, 8, g_bot, 13)
    rl = g_bot + 1
    leit_e = (f'=COUNTIF(tbEstoque[Status_Estoque],"REPOR")&" em REPOR ("&{KK}&" de criticidade Alta) · "&COUNTIF(tbEstoque[Status_Estoque],"ATENÇÃO")&" em ATENÇÃO ("'
              f'&COUNTIFS(tbEstoque[Criticidade],"Alta",tbEstoque[Status_Estoque],"ATENÇÃO")&" de criticidade Alta)"'
              f'&IF({A}B{pn["oculto"]}>0,CHAR(10)&"+"&{A}B{pn["oculto"]}&" item(ns) fora do painel: ver aba Analises, seção 7b","")')
    leit_r = ('="REPOR: atual <= mínimo · ATENÇÃO: até "&ROUND(pMargemAtencao*100,0)&"% acima do mínimo · Crítico em risco: criticidade Alta em REPOR (texto vermelho na fila)"')
    leitura(rl, 2, 7, leit_e, 30)
    leitura(rl, 8, 13, leit_r, 30)
    moldura(rt, 2, rl, 7)
    moldura(rt, 8, rl, 13)
    ws.row_dimensions[rl + 1].height = 8

    # ---- rodapé
    rf = rl + 2
    for k in range(3):
        ws.row_dimensions[rf + k].height = 16
    mescla(ws, rf, 2, 13, "Rastreabilidade: cada número deste painel é fórmula. DASHBOARD ← Analises ← Indicadores ← tbConsumo / tbEstoque / tbMateriais (CSVs em data/).",
           font=f(8, False, CINZA_T), ali=ali("left", "center"))
    mescla(ws, rf + 1, 2, 13,
           f'="Conferências da camada analítica: "&{A}B{cf["resumo"]}&" OK  ·  Qualidade dos dados: "&Qualidade!B6&" PASS / "&Qualidade!B7&" FAIL  ·  Conferências de Indicadores: "&Indicadores!B16&" OK"',
           font=f(8, True, CINZA_T), ali=ali("left", "center"))
    mescla(ws, rf + 2, 2, 13, "Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio. O projeto não realiza previsão e nenhum resultado é recomendação operacional real.",
           font=f(8, False, VERM_E, i=True), ali=ali("left", "center"))
    ultima = rf + 3

    # ---- página / vista
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = COBRE
    ws.print_area = f"A1:N{ultima}"
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", 9
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top, ws.page_margins.bottom = 0.3, 0.5
    ws.oddFooter.center.text = "Lobo Insights Industrial V0.2 · dados 100% sintéticos · página &P de &N"
    for q in estado["quebras"]:
        ws.row_breaks.append(Break(id=q))
    return {"ultima": ultima, "quebras": estado["quebras"], "estoque_lista": (l0, l1), "kpi_row": KPI_VALOR}


# ================================================================== ajustes nas demais abas
def ajustar_parametros(wb):
    """Acrescenta os 3 parâmetros da V0.2 abaixo do conteúdo da V0.1.1 (A20:C24); A1:F18 não é tocado."""
    ws = wb["Parametros"]
    ws["A20"] = "Parâmetros adicionados na V0.2 (entradas do autor)"
    ws["A20"].font = f(11, True, NAVY)
    for j, t_ in enumerate(["Parâmetro", "Valor", "Descrição"], start=1):
        c = ws.cell(21, j, t_)
        c.font, c.fill = f(10, True, "FFFFFF"), fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(*(lado("BFBFBF"),) * 4)
    novos = [
        ("Limite_Pareto", 0.8, "0%", "pLimitePareto", "Limite do % acumulado do Pareto (linha tracejada e 'Núcleo'). Convenção 80/20 do estudo; editável (0 a 1)."),
        ("Top_N_A", 3, "0", "pTopA", "Nº de maiores materiais somados no cartão de concentração A (1 a 20)."),
        ("Top_N_B", 5, "0", "pTopB", "Nº de maiores materiais somados no cartão de concentração B (1 a 20)."),
    ]
    for k, (nome, val, fmt, nd, desc) in enumerate(novos, start=22):
        ws.cell(k, 1, nome).font = f(10, True)
        c = ws.cell(k, 2, val)
        c.font, c.fill, c.number_format, c.border = f(10, cor="0000FF"), fill("FFFF00"), fmt, Border(*(lado("BFBFBF"),) * 4)
        ws.cell(k, 3, desc + " Nome definido: " + nd + ".").font = f(9, i=True, cor="595959")
        wb.defined_names[nd] = DefinedName(nd, attr_text="Parametros!$B$" + str(k))
    dv1 = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False, showErrorMessage=True,
                         errorTitle="Valor inválido", error="Informe um valor entre 0 e 1 (ex.: 0,8 = 80%).")
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="20", allow_blank=False, showErrorMessage=True,
                         errorTitle="Valor inválido", error="Informe um inteiro entre 1 e 20.")
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B22")
    dv2.add("B23:B24")


def ajustar_indicadores(wb):
    ws = wb["Indicadores"]
    ws["A1"] = "Indicadores – 6 KPIs oficiais e resumos (camada de indicadores; alimenta a aba Analises e o DASHBOARD)"
    ws["A3"] = ("Valores monetários em unidade fictícia. Todos os números são fórmulas sobre tbConsumo, tbEstoque e tbMateriais (nada digitado). "
                "Conteúdo idêntico ao da V0.1.1; a apresentação executiva está na aba DASHBOARD.")


def reescrever_leia_me(wb, meta, n_m, n_e, n_c):
    ws = wb["LEIA_ME"]
    ws.delete_rows(1, ws.max_row)
    ws["A1"] = "Lobo Insights Industrial – V0.2 Analytics & Dashboard"
    ws["A1"].font = f(14, True, NAVY)
    ws["A2"] = AVISO
    ws["A2"].font = f(11, True, "9C0006")
    linhas = [
        ("Pergunta central", "O que aconteceu na operação e o que precisa de atenção? Análise do PRESENTE e do histórico; este projeto NÃO faz previsão."),
        ("Versão / escopo", "V0.2 – camada analítica e DASHBOARD sobre a base V0.1.1 (dados, regras, 6 KPIs e fórmulas da V0.1.1 preservados). Sem IA, sem Power BI, sem n8n, sem previsão."),
        ("Fonte oficial dos números", "Os três CSVs da pasta data/. Materiais, Estoque e Consumo são cópias deles (tbMateriais, tbEstoque, tbConsumo). Nenhum número é digitado nas abas de resultado."),
        ("", ""),
        ("Camadas (do dado à imagem)", ""),
        ("Dados", "Materiais, Estoque, Consumo (tabelas estruturadas)."),
        ("Parâmetros", "Parametros: margem de ATENÇÃO, período, listas de valores válidos e (V0.2) limite do Pareto e nº de maiores materiais somados."),
        ("Validações", "Qualidade: 36 verificações ao vivo (Q01–Q36). Indicadores, seção 6, e Analises, seção 8: conferências de fechamento."),
        ("Indicadores", "Os 6 KPIs oficiais e resumos por mês, categoria, centro e material (V0.1.1, sem alteração de fórmulas)."),
        ("Cálculos", "Analises: evolução mensal, rankings por categoria/centro/material, Pareto de custo e de consumo, concentração, fila de prioridade de estoque, matriz criticidade × status e conferências."),
        ("Visualizações", "DASHBOARD: 6 cartões de KPI, 9 gráficos (mensal, categoria e centro em unidades e custo; 2 Paretos; estoque atual × mínimo), 6 cartões de concentração, fila de atenção de estoque e leituras geradas por fórmula."),
        ("", ""),
        ("Como ler o DASHBOARD", "De cima para baixo: KPIs → evolução mensal → categorias → centros → Pareto de materiais → estoque e criticidade. Os links da faixa superior levam às abas de apoio. Nenhum número é digitado: passe o cursor/selecione a célula para ver a fórmula."),
        ("Filtros / segmentações", "Não há filtros no DASHBOARD (decisão técnica registrada em docs/V0_2_NOTES.md): segmentações nativas exigem tabelas dinâmicas e não puderam ser verificadas neste ambiente; listas suspensas por fórmula fariam os painéis divergirem dos 6 KPIs oficiais. As tabelas de dados têm filtro de coluna."),
        ("Legenda de cores", "Cabeçalho azul-marinho = importado do CSV (ou cabeçalho de bloco). Cabeçalho verde = coluna calculada da V0.1.1. Célula amarela com texto azul = parâmetro editável (Parametros). Estoque: vermelho = REPOR, âmbar = ATENÇÃO, verde = OK."),
        ("Status de estoque", "REPOR: Estoque_Atual <= Estoque_Minimo. ATENÇÃO: Atual > Mínimo e <= Mínimo x (1 + Margem_ATENCAO; padrão 25%). OK: acima disso. É o estado ATUAL – não é previsão nem cálculo de reposição."),
        ("Critico_Em_Risco", "Sim quando Criticidade = Alta e Status_Estoque = REPOR."),
        ("Pareto", "Materiais ordenados do maior para o menor valor; % acumulado sobre o total oficial (KPI). 'Núcleo' = materiais necessários para atingir o limite (Parametros, padrão 80%). É descrição do período observado, não regra de gestão."),
        ("Custo estimado", "Custo_Movimentado = Quantidade x Custo_Unitario (2 casas). Custos e moeda são fictícios."),
        ("Como validar", "1) Qualidade: todos PASS. 2) Analises, seção 8, e Indicadores, seção 6: conferências OK (o rodapé do DASHBOARD resume). 3) Compare os KPIs com scripts/validar_dados.py e docs/TEST_REPORT.md."),
        ("Limitações", "Dados sintéticos; 3 meses; sem previsão; não testado no Microsoft Excel real (docs/LIMITATIONS.md). Diferenças entre meses vêm do sorteio dos dados, não de fenômeno operacional."),
        ("", ""),
        ("SHA-256 dos CSVs de origem", ""),
    ]
    for nome, h in (meta.get("sha") or {}).items():
        linhas.append((nome, h))
    linhas += [("", ""), ("Uso", "Somente estudo e portfólio. Nenhum uso operacional real. Veja docs/LIMITATIONS.md.")]
    for k, (a, b) in enumerate(linhas, start=4):
        ca, cb = ws.cell(k, 1, a or None), ws.cell(k, 2, b or None)
        ca.font = f(10, True) if a else f(10)
        cb.font = f(10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
        if a in ("Camadas (do dado à imagem)", "SHA-256 dos CSVs de origem"):
            ca.font = f(11, True, NAVY)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120


# ================================================================== orquestração
def construir_v0_2(dfs, destino, meta=None):
    """Gera o XLSX cru da V0.2 (sem cache de valores). Retorna informações de posição das camadas novas."""
    meta = meta or {}
    tmp = Path(tempfile.mkdtemp(prefix="lobo_v02_"))
    base = tmp / "base_v0_1_1.xlsx"
    ind = construir(dfs, base, meta)                           # abas de dados, Qualidade, Indicadores, Parametros (V0.1.1)
    wb = load_workbook(base)
    n_m, n_e, n_c = len(dfs["mat"]), len(dfs["est"]), len(dfs["con"])
    ajustar_parametros(wb)
    ajustar_indicadores(wb)
    reescrever_leia_me(wb, meta, n_m, n_e, n_c)
    ws_a, P, r_conf = construir_analises(wb, ind, n_m)
    P = conferencias(ws_a, P, ind, r_conf, n_m)
    ws_d = wb.create_sheet(S_DASH)
    info_d = construir_dashboard(wb, ws_d, P, ind)

    ordem = [S_DASH, "LEIA_ME", "Indicadores", S_ANA, "Materiais", "Estoque", "Consumo", "Qualidade", "Parametros"]
    wb._sheets = [wb[n] for n in ordem]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == S_DASH
    ws_a.page_setup.orientation, ws_a.page_setup.fitToWidth, ws_a.page_setup.fitToHeight = "landscape", 1, 0
    ws_a.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.properties.title = "Lobo Insights Industrial V0.2 - Analytics & Dashboard"
    wb.properties.creator = "Lobo Insights Industrial (dados sintéticos)"
    wb.calculation.fullCalcOnLoad = True
    wb.save(destino)
    shutil.rmtree(tmp, ignore_errors=True)
    return {"ind": ind, "P": P, "dash": info_d}


def main():
    from validar_dados import carregar
    dfs = carregar()
    meta = {"sha": {p.name: sha256(p) for p in sorted(DATA.glob("*.csv"))}}
    tmp = Path(tempfile.mkdtemp(prefix="lobo_build_v02_"))
    cru = tmp / "cru.xlsx"
    construir_v0_2(dfs, cru, meta)
    try:
        valores, info = recalcular(cru)
    except LibreOfficeIndisponivel as e:
        raise SystemExit(f"ERRO: {e}")
    print("recalc:", info)
    if info.get("status") != "success" or info.get("total_errors"):
        raise SystemExit("Erros de fórmula após o recálculo – corrija antes de continuar.")
    EXCEL.mkdir(exist_ok=True)
    destino = EXCEL / NOME_XLSX_V02
    injetar_cache(cru, valores, destino)
    shutil.rmtree(tmp, ignore_errors=True)
    print("gerado:", destino, "sha256:", sha256(destino))


if __name__ == "__main__":
    main()
