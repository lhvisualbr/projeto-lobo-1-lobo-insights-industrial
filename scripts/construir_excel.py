# -*- coding: utf-8 -*-
"""Constrói excel/lobo_insights_industrial_v0_1.xlsx a partir dos CSVs.

Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio.

Etapas:
 1. openpyxl escreve tabelas estruturadas, fórmulas reais, validação de dados e
    formatação condicional (arquivo "cru", sem valores em cache);
 2. o LibreOffice (modo headless) recalcula uma CÓPIA e os valores calculados são lidos;
 3. os valores são gravados como cache dentro do arquivo cru (as fórmulas continuam
    fórmulas). Assim, visualizadores, pandas e celulares mostram os números, e o Excel
    ainda recalcula tudo ao abrir (fullCalcOnLoad).

Requer LibreOffice apenas para o passo 2. O executável é procurado pela variável de ambiente
LOBO_SOFFICE (caminho do `soffice`) ou, se ela não existir, no PATH (`soffice`/`libreoffice`).
Sem LibreOffice o arquivo cru continua válido no Excel, e a suíte de testes marca como SKIP
(nunca como PASS) os testes que dependem do recálculo.
"""
import datetime as dt
import hashlib
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from lobo_common import (CATEGORIAS, CENTROS, CHECKS, CRITICIDADES, EXCEL, MARGEM_ATENCAO,
                         PERIODO_FIM, PERIODO_INICIO, TIPOS_MOV, UNIDADES)

NOME_XLSX = "lobo_insights_industrial_v0_1.xlsx"
AVISO = "Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio."

# ------------------------------------------------------------------ estilos
FONTE = "Arial"
NAVY, VERDE, CINZA = "1F3A5F", "2E7D5B", "F2F2F2"
f_norm = Font(name=FONTE, size=10)
f_bold = Font(name=FONTE, size=10, bold=True)
f_hdr = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
f_tit = Font(name=FONTE, size=14, bold=True, color=NAVY)
f_sec = Font(name=FONTE, size=11, bold=True, color=NAVY)
f_nota = Font(name=FONTE, size=9, italic=True, color="595959")
f_input = Font(name=FONTE, size=10, color="0000FF")
fill_hdr = PatternFill("solid", fgColor=NAVY)
fill_calc = PatternFill("solid", fgColor=VERDE)
fill_cinza = PatternFill("solid", fgColor=CINZA)
fill_input = PatternFill("solid", fgColor="FFFF00")
fill_red = PatternFill("solid", bgColor="F8CBCB", fgColor="F8CBCB")
fill_amb = PatternFill("solid", bgColor="FFE9A8", fgColor="FFE9A8")
fill_ok = PatternFill("solid", bgColor="CDEBD6", fgColor="CDEBD6")
borda = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

NUMERICAS = {"Custo_Unitario", "Estoque_Atual", "Estoque_Minimo", "Lead_Time_Dias", "Quantidade"}


def _conv(col, s):
    """Converte o texto do CSV para o tipo Excel; valores estranhos permanecem como texto."""
    if s is None or str(s).strip() == "":
        return None
    s = str(s)
    if col == "Data":
        try:
            return dt.date.fromisoformat(s)
        except ValueError:
            return s
    if col in NUMERICAS:
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d+\.\d+", s):
            return float(s)
    return s


def _tabela(ws, nome, df, calc, comentarios=None):
    """Escreve df + colunas calculadas e cria a tabela estruturada. Retorna (primeira, ultima_linha)."""
    cols = list(df.columns) + [c for c, _ in calc]
    ws.append(cols)
    for _, linha in df.iterrows():
        vals = [_conv(c, linha[c]) for c in df.columns]
        ws.append(vals + [f.replace("{T}", nome) for _, f in calc])
    n = len(df) + 1
    ref = f"A1:{get_column_letter(len(cols))}{n}"
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    for j, c in enumerate(cols, start=1):
        cel = ws.cell(row=1, column=j)
        cel.font, cel.alignment = f_hdr, Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.fill = fill_calc if j > len(df.columns) else fill_hdr
        if comentarios and c in comentarios:
            cel.comment = Comment(comentarios[c], "Lobo Insights")
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cel in row:
            cel.font = f_norm
    for j, c in enumerate(cols, start=1):
        amostra = [str(ws.cell(row=r, column=j).value or "") for r in range(2, min(n, 60) + 1)]
        larg = max([len(str(c))] + [len(a) for a in amostra if not a.startswith("=")])
        ws.column_dimensions[get_column_letter(j)].width = min(max(larg + 3, 14), 46)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    return 2, n


def _dv_lista(ws, faixa, nome_lista, titulo):
    dv = DataValidation(type="list", formula1=f"={nome_lista}", allow_blank=False,
                        showErrorMessage=True, errorTitle="Valor inválido",
                        error=f"Use um valor da lista {titulo} (aba Parametros).")
    ws.add_data_validation(dv)
    dv.add(faixa)


def _dv_num(ws, faixa, tipo, op, f1, msg):
    dv = DataValidation(type=tipo, operator=op, formula1=f1, allow_blank=False,
                        showErrorMessage=True, errorTitle="Valor inválido", error=msg)
    ws.add_data_validation(dv)
    dv.add(faixa)


def _cabecalho_bloco(ws, linha, textos, colunas_ini=1):
    for j, t in enumerate(textos, start=colunas_ini):
        c = ws.cell(row=linha, column=j, value=t)
        c.font, c.fill, c.border = f_hdr, fill_hdr, borda
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _titulo(ws, texto):
    ws["A1"] = texto
    ws["A1"].font = f_tit
    ws["A2"] = AVISO
    ws["A2"].font = f_nota


# ------------------------------------------------------------------ construção
def construir(dfs, destino, meta=None):
    meta = meta or {}
    mat, est, con = dfs["mat"], dfs["est"], dfs["con"]
    wb = Workbook()
    wb.properties.creator = "Lobo Insights Industrial (dados sintéticos)"
    wb.properties.title = "Lobo Insights Industrial V0.1"
    wb.calculation.fullCalcOnLoad = True

    ws_leia = wb.active
    ws_leia.title = "LEIA_ME"
    ws_m = wb.create_sheet("Materiais")
    ws_e = wb.create_sheet("Estoque")
    ws_c = wb.create_sheet("Consumo")
    ws_q = wb.create_sheet("Qualidade")
    ws_i = wb.create_sheet("Indicadores")
    ws_p = wb.create_sheet("Parametros")

    # ---------------- Parametros (parâmetros e listas de valores válidos)
    _titulo(ws_p, "Parâmetros e listas de valores válidos")
    _cabecalho_bloco(ws_p, 4, ["Parâmetro", "Valor", "Descrição"])
    params = [
        ("Margem_ATENCAO", MARGEM_ATENCAO, "0%", "ATENÇÃO quando Estoque_Atual > Mínimo e <= Mínimo x (1 + margem). Nome definido: pMargemAtencao. Valor de estudo (entrada do autor)."),
        ("Inicio_Periodo", dt.date.fromisoformat(PERIODO_INICIO), "yyyy-mm-dd", "Primeiro dia do período de consumo (3 meses consecutivos, conforme o PDF). Nome: pInicioPeriodo."),
        ("Fim_Periodo", dt.date.fromisoformat(PERIODO_FIM), "yyyy-mm-dd", "Último dia do período. Nome: pFimPeriodo."),
    ]
    for k, (nome, val, fmt, desc) in enumerate(params, start=5):
        ws_p.cell(row=k, column=1, value=nome).font = f_bold
        c = ws_p.cell(row=k, column=2, value=val)
        c.font, c.fill, c.number_format, c.border = f_input, fill_input, fmt, borda
        ws_p.cell(row=k, column=3, value=desc).font = f_nota
    listas = [("Categorias", CATEGORIAS, "lstCategorias"), ("Unidades", UNIDADES, "lstUnidades"),
              ("Criticidade", CRITICIDADES, "lstCriticidade"), ("Centros de trabalho", CENTROS, "lstCentros"),
              ("Tipos de movimentação", TIPOS_MOV, "lstTipoMov")]
    _cabecalho_bloco(ws_p, 10, [l[0] for l in listas] + ["Meses do período (fórmula)"])
    for j, (_, itens, nome) in enumerate(listas, start=1):
        for k, v in enumerate(itens, start=11):
            ws_p.cell(row=k, column=j, value=v).font = f_norm
        L = get_column_letter(j)
        wb.defined_names[nome] = DefinedName(nome, attr_text=f"Parametros!${L}$11:${L}${10 + len(itens)}")
    for k in range(3):
        ws_p.cell(row=11 + k, column=6, value=(
            f'=YEAR(EDATE(pInicioPeriodo,{k}))&"-"&TEXT(MONTH(EDATE(pInicioPeriodo,{k})),"00")')).font = f_norm
    wb.defined_names["lstMeses"] = DefinedName("lstMeses", attr_text="Parametros!$F$11:$F$13")
    wb.defined_names["pMargemAtencao"] = DefinedName("pMargemAtencao", attr_text="Parametros!$B$5")
    wb.defined_names["pInicioPeriodo"] = DefinedName("pInicioPeriodo", attr_text="Parametros!$B$6")
    wb.defined_names["pFimPeriodo"] = DefinedName("pFimPeriodo", attr_text="Parametros!$B$7")
    for j, w in enumerate([24, 16, 16, 24, 24, 28], start=1):
        ws_p.column_dimensions[get_column_letter(j)].width = w
    ws_p["A18"] = ("Os valores acima são as listas oficiais usadas pela validação de dados e pelas verificações da aba Qualidade. "
                   "Alterar uma lista muda o que é considerado válido.")
    ws_p["A18"].font = f_nota

    # ---------------- Materiais
    _, n_m = _tabela(ws_m, "tbMateriais", mat, [])
    _dv_lista(ws_m, f"C2:C{n_m}", "lstCategorias", "Categorias")
    _dv_lista(ws_m, f"D2:D{n_m}", "lstUnidades", "Unidades")
    _dv_lista(ws_m, f"F2:F{n_m}", "lstCriticidade", "Criticidade")
    _dv_num(ws_m, f"E2:E{n_m}", "decimal", "greaterThan", "0", "Custo_Unitario deve ser maior que zero.")
    ws_m.conditional_formatting.add(f"F2:F{n_m}", CellIsRule(operator="equal", formula=['"Alta"'], font=Font(name=FONTE, bold=True, color="9C0006")))
    for r in range(2, n_m + 1):
        ws_m.cell(row=r, column=5).number_format = "#,##0.00"
        ws_m.cell(row=r, column=5).alignment = Alignment(horizontal="center")

    # ---------------- Estoque
    T = "{T}"
    calc_e = [
        ("Status_Estoque", f'=IF({T}[[#This Row],[Estoque_Atual]]<={T}[[#This Row],[Estoque_Minimo]],"REPOR",'
                           f'IF({T}[[#This Row],[Estoque_Atual]]<={T}[[#This Row],[Estoque_Minimo]]*(1+pMargemAtencao),"ATENÇÃO","OK"))'),
        ("Critico_Em_Risco", f'=IF(AND({T}[[#This Row],[Criticidade]]="Alta",{T}[[#This Row],[Status_Estoque]]="REPOR"),"Sim","Não")'),
    ]
    coment_e = {
        "Status_Estoque": "REPOR se Estoque_Atual <= Estoque_Minimo. ATENÇÃO se Atual > Mínimo e Atual <= Mínimo x (1 + Margem_ATENCAO). Caso contrário OK. Regra de estado ATUAL, não é previsão.",
        "Critico_Em_Risco": "Sim quando Criticidade = Alta e Status_Estoque = REPOR (Estoque_Atual <= Mínimo). Definição conservadora: só REPOR, não inclui ATENÇÃO.",
    }
    _, n_e = _tabela(ws_e, "tbEstoque", est, calc_e, coment_e)
    _dv_lista(ws_e, f"C2:C{n_e}", "lstCategorias", "Categorias")
    _dv_lista(ws_e, f"D2:D{n_e}", "lstUnidades", "Unidades")
    _dv_lista(ws_e, f"J2:J{n_e}", "lstCriticidade", "Criticidade")
    _dv_num(ws_e, f"E2:F{n_e}", "whole", "greaterThanOrEqual", "0", "Estoque deve ser inteiro >= 0.")
    _dv_num(ws_e, f"G2:G{n_e}", "decimal", "greaterThan", "0", "Custo_Unitario deve ser maior que zero.")
    _dv_num(ws_e, f"H2:H{n_e}", "whole", "greaterThan", "0", "Lead_Time_Dias deve ser inteiro > 0.")
    ws_e.conditional_formatting.add(f"A2:L{n_e}", FormulaRule(formula=['$K2="REPOR"'], fill=fill_red))
    ws_e.conditional_formatting.add(f"A2:L{n_e}", FormulaRule(formula=['$K2="ATENÇÃO"'], fill=fill_amb))
    ws_e.conditional_formatting.add(f"L2:L{n_e}", CellIsRule(operator="equal", formula=['"Sim"'], font=Font(name=FONTE, bold=True, color="9C0006")))
    ws_e.conditional_formatting.add(f"J2:J{n_e}", CellIsRule(operator="equal", formula=['"Alta"'], font=Font(name=FONTE, bold=True, color="9C0006")))
    for r in range(2, n_e + 1):
        ws_e.cell(row=r, column=7).number_format = "#,##0.00"
        for j in (5, 6, 8):
            ws_e.cell(row=r, column=j).alignment = Alignment(horizontal="center")

    # ---------------- Consumo
    calc_c = [
        ("Custo_Movimentado", f"=ROUND({T}[[#This Row],[Quantidade]]*{T}[[#This Row],[Custo_Unitario]],2)"),
        ("Mes", f'=YEAR({T}[[#This Row],[Data]])&"-"&TEXT(MONTH({T}[[#This Row],[Data]]),"00")'),
    ]
    coment_c = {
        "Custo_Movimentado": "Quantidade x Custo_Unitario, arredondado a 2 casas (custo estimado, valor fictício).",
        "Mes": "Ano-mês (aaaa-mm) da Data. Usa TEXT(...,\"00\") para não depender do idioma do Excel.",
    }
    _, n_c = _tabela(ws_c, "tbConsumo", con, calc_c, coment_c)
    _dv_num(ws_c, f"B2:B{n_c}", "date", "between", "pInicioPeriodo", "Data fora do período ou inválida.")
    ws_c.data_validations.dataValidation[-1].formula2 = "pFimPeriodo"
    _dv_lista(ws_c, f"E2:E{n_c}", "lstCategorias", "Categorias")
    _dv_lista(ws_c, f"G2:G{n_c}", "lstUnidades", "Unidades")
    _dv_lista(ws_c, f"I2:I{n_c}", "lstCentros", "Centros de trabalho")
    _dv_lista(ws_c, f"J2:J{n_c}", "lstTipoMov", "Tipos de movimentação")
    _dv_num(ws_c, f"F2:F{n_c}", "whole", "greaterThan", "0", "Quantidade deve ser inteiro > 0.")
    _dv_num(ws_c, f"H2:H{n_c}", "decimal", "greaterThan", "0", "Custo_Unitario deve ser maior que zero.")
    dv_id = DataValidation(type="list", formula1=f"=Materiais!$A$2:$A${n_m}", allow_blank=False,
                           showErrorMessage=True, errorTitle="Material inexistente",
                           error="Material_ID precisa existir na aba Materiais.")
    ws_c.add_data_validation(dv_id)
    dv_id.add(f"C2:C{n_c}")
    for r in range(2, n_c + 1):
        ws_c.cell(row=r, column=2).number_format = "yyyy-mm-dd"
        ws_c.cell(row=r, column=8).number_format = "#,##0.00"
        ws_c.cell(row=r, column=11).number_format = "#,##0.00"
        for j in (6, 8, 11, 12):
            ws_c.cell(row=r, column=j).alignment = Alignment(horizontal="center")

    # ---------------- Qualidade
    _titulo(ws_q, "Qualidade dos dados – verificações ao vivo (fórmulas)")
    ws_q["A3"] = ("Cada linha é uma fórmula sobre tbMateriais, tbEstoque e tbConsumo. Erros = 0 => PASS. "
                  "O mesmo conjunto (Q01–Q36) é executado de forma independente em Python (scripts/validar_dados.py).")
    ws_q["A3"].font = f_nota
    q0, q1 = 10, 10 + len(CHECKS) - 1
    for k, (rot, fml) in enumerate([("Testes executados", f"=COUNTA(A{q0}:A{q1})"),
                                    ("PASS", f'=COUNTIF(E{q0}:E{q1},"PASS")'),
                                    ("FAIL", f'=COUNTIF(E{q0}:E{q1},"FAIL")')], start=5):
        ws_q.cell(row=k, column=1, value=rot).font = f_bold
        c = ws_q.cell(row=k, column=2, value=fml)
        c.font, c.alignment = f_bold, Alignment(horizontal="left")
    ws_q.append([])
    for j, t in enumerate(["Teste", "Regra", "Resultado", "Quantidade de erros", "Status"], start=1):
        ws_q.cell(row=9, column=j, value=t)
    for k, (cid, teste, regra, fml) in enumerate(CHECKS, start=q0):
        ws_q.cell(row=k, column=1, value=f"{cid} · {teste}")
        ws_q.cell(row=k, column=2, value=regra)
        ws_q.cell(row=k, column=3, value=f'=IF(D{k}=0,"Nenhum erro encontrado","Erros encontrados")')
        ws_q.cell(row=k, column=4, value=fml)
        ws_q.cell(row=k, column=5, value=f'=IF(D{k}=0,"PASS","FAIL")')
    tq = Table(displayName="tbQualidade", ref=f"A9:E{q1}")
    tq.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws_q.add_table(tq)
    for j in range(1, 6):
        c = ws_q.cell(row=9, column=j)
        c.font, c.fill, c.alignment = f_hdr, fill_hdr, Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in range(q0, q1 + 1):
            ws_q.cell(row=r, column=j).font = f_norm
    for j, w in enumerate([44, 78, 26, 20, 10], start=1):
        ws_q.column_dimensions[get_column_letter(j)].width = w
    ws_q.freeze_panes = "A10"
    ws_q.conditional_formatting.add(f"E{q0}:E{q1}", CellIsRule(operator="equal", formula=['"PASS"'], fill=fill_ok, font=Font(name=FONTE, bold=True, color="1E6B3A")))
    ws_q.conditional_formatting.add(f"E{q0}:E{q1}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=fill_red, font=Font(name=FONTE, bold=True, color="9C0006")))
    for r in range(q0, q1 + 1):
        ws_q.cell(row=r, column=4).alignment = Alignment(horizontal="center")
        ws_q.cell(row=r, column=5).alignment = Alignment(horizontal="center")

    # ---------------- Indicadores
    _titulo(ws_i, "Indicadores V0.1 – Lobo Insights Industrial")
    ws_i["A3"] = ("Valores monetários em unidade fictícia. Todos os números são fórmulas sobre tbConsumo, tbEstoque e tbMateriais "
                  "(nada digitado). Este NÃO é o painel final (V0.2).")
    ws_i["A3"].font = f_nota
    n_meses, n_cat, n_cen, n_mat = 3, len(CATEGORIAS), len(CENTROS), len(mat)
    r_kpi = 6                                   # 6 KPIs em 6..11
    r_info = 13                                 # informativos 13..14
    r_conf = 16                                 # resumo das conferências
    s_mes = 19; s_cat = s_mes + n_meses + 4; s_cen = s_cat + n_cat + 4; s_mat = s_cen + n_cen + 4
    tot = {}                                    # linha de total por bloco
    ws_i.cell(row=4, column=1, value="1. KPIs principais").font = f_sec
    _cabecalho_bloco(ws_i, 5, ["KPI", "Valor", "Como é calculado", "", "", "", "Origem"])
    ws_i.merge_cells("C5:F5"); ws_i.merge_cells("G5:J5")
    m0, m1 = s_mat + 2, s_mat + 1 + n_mat       # linhas de dados do resumo por material
    kpis = [
        ("Unidades consumidas", '=SUMIFS(tbConsumo[Quantidade],tbConsumo[Tipo_Movimentacao],"Consumo")', "#,##0",
         "Soma de Quantidade das movimentações do tipo Consumo (SOMASES).", "tbConsumo[Quantidade]"),
        ("Custo estimado consumido", '=SUMIFS(tbConsumo[Custo_Movimentado],tbConsumo[Tipo_Movimentacao],"Consumo")', "#,##0.00",
         "Soma de Custo_Movimentado (Quantidade x Custo_Unitario) das movimentações de Consumo (SOMASES). Valor fictício.", "tbConsumo[Custo_Movimentado]"),
        ("Movimentações", '=COUNTIFS(tbConsumo[Tipo_Movimentacao],"Consumo")', "#,##0",
         "Quantidade de registros (linhas) do tipo Consumo (CONT.SES). Não é quantidade de unidades.", "tbConsumo[Movimento_ID]"),
        ("Materiais distintos movimentados", f'=COUNTIF(B{m0}:B{m1},">0")', "#,##0",
         "Materiais do catálogo com pelo menos 1 movimentação de Consumo (contagem sobre o resumo por material, seção 5).", "Seção 5 desta aba"),
        ("Itens abaixo ou iguais ao mínimo", '=COUNTIF(tbEstoque[Status_Estoque],"REPOR")', "#,##0",
         "Itens com Status_Estoque = REPOR (Estoque_Atual <= Estoque_Minimo).", "tbEstoque[Status_Estoque]"),
        ("Itens críticos abaixo ou iguais ao mínimo", '=COUNTIF(tbEstoque[Critico_Em_Risco],"Sim")', "#,##0",
         "Itens com Criticidade = Alta e Estoque_Atual <= Estoque_Minimo (Critico_Em_Risco = Sim).", "tbEstoque[Critico_Em_Risco]"),
    ]
    for k, (rot, fml, fmt, desc, orig) in enumerate(kpis):
        r = r_kpi + k
        ws_i.cell(row=r, column=1, value=rot).font = f_bold
        c = ws_i.cell(row=r, column=2, value=fml)
        c.font, c.number_format, c.alignment = Font(name=FONTE, size=11, bold=True, color=NAVY), fmt, Alignment(horizontal="right", vertical="center")
        ws_i.cell(row=r, column=1).alignment = Alignment(vertical="center")
        ws_i.cell(row=r, column=3, value=desc).font = f_norm
        ws_i.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="center")
        ws_i.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
        ws_i.cell(row=r, column=7, value=orig).font = f_nota
        ws_i.cell(row=r, column=7).alignment = Alignment(vertical="center")
        ws_i.merge_cells(start_row=r, start_column=7, end_row=r, end_column=10)
        ws_i.row_dimensions[r].height = 28
        for j in (1, 2):
            ws_i.cell(row=r, column=j).border = borda
    infos = [("Informativo: itens em ATENÇÃO", '=COUNTIF(tbEstoque[Status_Estoque],"ATENÇÃO")',
              "Atual > Mínimo e Atual <= Mínimo x (1 + margem). Não faz parte dos 6 KPIs."),
             ("Informativo: itens de criticidade Alta em ATENÇÃO", '=COUNTIFS(tbEstoque[Criticidade],"Alta",tbEstoque[Status_Estoque],"ATENÇÃO")',
              "Complementa o KPI 6, que considera apenas REPOR. Não faz parte dos 6 KPIs.")]
    for k, (rot, fml, desc) in enumerate(infos):
        r = r_info + k
        ws_i.cell(row=r, column=1, value=rot).font = f_norm
        c = ws_i.cell(row=r, column=2, value=fml); c.font = f_norm; c.alignment = Alignment(horizontal="right")
        ws_i.cell(row=r, column=3, value=desc).font = f_nota
        ws_i.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)

    def bloco(ini, titulo, rotulo, chaves_fml, campo_tbl, n):
        """Resumo com Movimentações, Unidades, Custo e % do custo."""
        ws_i.cell(row=ini, column=1, value=titulo).font = f_sec
        _cabecalho_bloco(ws_i, ini + 1, [rotulo, "Movimentações", "Unidades consumidas", "Custo estimado", "% do custo"])
        r0, r1, rt = ini + 2, ini + 1 + n, ini + 2 + n
        for k in range(n):
            r = r0 + k
            ws_i.cell(row=r, column=1, value=chaves_fml(k)).font = f_norm
            ws_i.cell(row=r, column=2, value=f'=COUNTIFS(tbConsumo[{campo_tbl}],$A{r},tbConsumo[Tipo_Movimentacao],"Consumo")')
            ws_i.cell(row=r, column=3, value=f'=SUMIFS(tbConsumo[Quantidade],tbConsumo[{campo_tbl}],$A{r},tbConsumo[Tipo_Movimentacao],"Consumo")')
            ws_i.cell(row=r, column=4, value=f'=SUMIFS(tbConsumo[Custo_Movimentado],tbConsumo[{campo_tbl}],$A{r},tbConsumo[Tipo_Movimentacao],"Consumo")')
            ws_i.cell(row=r, column=5, value=f"=IF(D${rt}=0,0,D{r}/D${rt})")
        ws_i.cell(row=rt, column=1, value="Total").font = f_bold
        for j, L in ((2, "B"), (3, "C"), (4, "D"), (5, "E")):
            ws_i.cell(row=rt, column=j, value=f"=SUM({L}{r0}:{L}{r1})")
        for r in range(r0, rt + 1):
            for j, fmt in ((2, "#,##0"), (3, "#,##0"), (4, "#,##0.00"), (5, "0.0%")):
                c = ws_i.cell(row=r, column=j); c.number_format = fmt; c.font = f_bold if r == rt else f_norm
            for j in range(1, 6):
                ws_i.cell(row=r, column=j).border = borda
                if r == rt:
                    ws_i.cell(row=r, column=j).fill = fill_cinza
        return rt

    tot["mes"] = bloco(s_mes, "2. Resumo por mês", "Mês (aaaa-mm)",
                       lambda k: f"=Parametros!$F${11 + k}", "Mes", n_meses)
    tot["cat"] = bloco(s_cat, "3. Resumo por categoria", "Categoria",
                       lambda k: f"=Parametros!$A${11 + k}", "Categoria", n_cat)
    tot["cen"] = bloco(s_cen, "4. Resumo por centro de trabalho", "Centro_Trabalho",
                       lambda k: f"=Parametros!$D${11 + k}", "Centro_Trabalho", n_cen)
    tot["mat"] = bloco(s_mat, "5. Resumo por material", "Material_ID",
                       lambda k: f"=Materiais!$A${2 + k}", "Material_ID", n_mat)
    # colunas descritivas do resumo por material (procura por chave: INDEX/MATCH, alternativa compatível ao PROCX)
    for j, t in enumerate(["Material", "Categoria", "Criticidade", "Estoque_Atual", "Status_Estoque"], start=6):
        c = ws_i.cell(row=s_mat + 1, column=j, value=t)
        c.font, c.fill, c.border = f_hdr, fill_calc, borda
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for k in range(n_mat):
        r = m0 + k
        for j, (tbl, campo, alt) in enumerate([("tbMateriais", "Material", "Não encontrado"), ("tbMateriais", "Categoria", "Não encontrado"),
                                               ("tbMateriais", "Criticidade", "Não encontrado"), ("tbEstoque", "Estoque_Atual", "Sem estoque"),
                                               ("tbEstoque", "Status_Estoque", "Sem estoque")], start=6):
            c = ws_i.cell(row=r, column=j, value=f'=IFERROR(INDEX({tbl}[{campo}],MATCH($A{r},{tbl}[Material_ID],0)),"{alt}")')
            c.font, c.border = f_norm, borda
    ws_i.conditional_formatting.add(f"J{m0}:J{m1}", CellIsRule(operator="equal", formula=['"REPOR"'], fill=fill_red, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws_i.conditional_formatting.add(f"J{m0}:J{m1}", CellIsRule(operator="equal", formula=['"ATENÇÃO"'], fill=fill_amb, font=Font(name=FONTE, bold=True, color="7F6000")))
    ws_i.conditional_formatting.add(f"H{m0}:H{m1}", CellIsRule(operator="equal", formula=['"Alta"'], font=Font(name=FONTE, bold=True, color="9C0006")))

    # conferências internas: total de cada resumo == KPI correspondente
    s_conf = s_mat + n_mat + 5
    ws_i.cell(row=s_conf, column=1, value="6. Conferências internas (cada resumo deve fechar com o KPI)").font = f_sec
    _cabecalho_bloco(ws_i, s_conf + 1, ["Conferência", "Total do resumo", "KPI", "Diferença", "Status"])
    kpi_ref = {"B": f"$B${r_kpi + 2}", "C": f"$B${r_kpi}", "D": f"$B${r_kpi + 1}"}
    rotulo_col = {"B": "Movimentações", "C": "Unidades", "D": "Custo estimado"}
    rc = s_conf + 2
    for bl, nome in (("mes", "mês"), ("cat", "categoria"), ("cen", "centro"), ("mat", "material")):
        for col in ("B", "C", "D"):
            ws_i.cell(row=rc, column=1, value=f"Soma por {nome} = {rotulo_col[col]}")
            ws_i.cell(row=rc, column=2, value=f"={col}{tot[bl]}")
            ws_i.cell(row=rc, column=3, value=f"={kpi_ref[col]}")
            ws_i.cell(row=rc, column=4, value=f"=B{rc}-C{rc}")
            ws_i.cell(row=rc, column=5, value=f'=IF(ABS(D{rc})<0.005,"OK","DIVERGE")')
            for j in range(1, 6):
                c = ws_i.cell(row=rc, column=j); c.font = f_norm; c.border = borda
            ws_i.cell(row=rc, column=2).number_format = "#,##0.00"
            ws_i.cell(row=rc, column=3).number_format = "#,##0.00"
            ws_i.cell(row=rc, column=4).number_format = "#,##0.00"
            rc += 1
    c0, c1 = s_conf + 2, rc - 1
    ws_i.conditional_formatting.add(f"E{c0}:E{c1}", CellIsRule(operator="equal", formula=['"OK"'], fill=fill_ok, font=Font(name=FONTE, bold=True, color="1E6B3A")))
    ws_i.conditional_formatting.add(f"E{c0}:E{c1}", CellIsRule(operator="equal", formula=['"DIVERGE"'], fill=fill_red, font=Font(name=FONTE, bold=True, color="9C0006")))
    ws_i.cell(row=r_conf, column=1, value="Conferências internas OK").font = f_bold
    c = ws_i.cell(row=r_conf, column=2, value=f'=COUNTIF(E{c0}:E{c1},"OK")&" de "&ROWS(E{c0}:E{c1})')
    c.font, c.alignment = f_bold, Alignment(horizontal="right")
    for j, w in enumerate([44, 16, 20, 18, 14, 30, 16, 14, 15, 15], start=1):
        ws_i.column_dimensions[get_column_letter(j)].width = w
    ws_i.freeze_panes = "A4"

    # ---------------- LEIA_ME
    _leia_me(ws_leia, meta, len(mat), len(est), len(con))

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title in ("Materiais", "Estoque", "Consumo", "Qualidade")
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    wb.save(destino)
    return {"kpi_row": r_kpi, "tot": tot, "s_conf": s_conf, "conf_rows": (c0, c1),
            "mat_rows": (m0, m1), "q_rows": (q0, q1)}


def _leia_me(ws, meta, n_m, n_e, n_c):
    ws["A1"] = "Lobo Insights Industrial – V0.1 Foundation (dados + Excel base)"
    ws["A1"].font = f_tit
    ws["A2"] = AVISO
    ws["A2"].font = Font(name=FONTE, size=11, bold=True, color="9C0006")
    linhas = [
        ("Pergunta central", "O que aconteceu na operação e o que precisa de atenção? (análise do PRESENTE e do histórico; não há previsão de demanda neste projeto)."),
        ("Versão / escopo", "V0.1 – dados sintéticos, tabelas estruturadas, verificações de qualidade e indicadores básicos. Sem dashboard (V0.2), sem IA, sem Power BI, sem n8n."),
        ("Fonte oficial dos números", "Os três CSVs da pasta data/. As abas Materiais, Estoque e Consumo são cópias desses CSVs (tbMateriais, tbEstoque, tbConsumo). Nenhum número é digitado à mão nas abas de resultado."),
        ("", ""),
        ("Abas", ""),
        ("Materiais", f"Catálogo fictício ({n_m} materiais, 6 categorias). Tabela tbMateriais."),
        ("Estoque", f"Posição de estoque fictícia ({n_e} itens) + colunas calculadas Status_Estoque e Critico_Em_Risco. Tabela tbEstoque."),
        ("Consumo", f"Movimentações fictícias ({n_c}), 3 meses consecutivos + colunas calculadas Custo_Movimentado e Mes. Tabela tbConsumo."),
        ("Qualidade", "36 verificações (Q01–Q36) em fórmulas ao vivo: teste, regra, resultado, quantidade de erros e status."),
        ("Indicadores", "6 KPIs, resumos por mês, categoria, centro e material, e conferências internas (cada resumo precisa fechar com o KPI)."),
        ("Parametros", "Margem de ATENÇÃO, período de análise e listas de valores válidos (usadas pela validação de dados)."),
        ("", ""),
        ("Legenda de cores", "Cabeçalho azul-marinho = coluna importada do CSV. Cabeçalho verde = coluna calculada por fórmula. Célula amarela com texto azul = parâmetro editável (aba Parametros)."),
        ("Status de estoque", "REPOR: Estoque_Atual <= Estoque_Minimo. ATENÇÃO: Estoque_Atual > Mínimo e <= Mínimo x (1 + Margem_ATENCAO; padrão 25%, editável na aba Parametros). OK: acima disso. É apenas o estado ATUAL – não é previsão nem cálculo de reposição."),
        ("Critico_Em_Risco", "Sim quando Criticidade = Alta e Status_Estoque = REPOR."),
        ("Custo estimado", "Custo_Movimentado = Quantidade x Custo_Unitario (arredondado a 2 casas). Custos e moeda são fictícios."),
        ("Funções usadas", "SOMASES (SUMIFS), CONT.SES (COUNTIFS), SE (IF), SEERRO (IFERROR), ÍNDICE+CORRESP (INDEX+MATCH, alternativa compatível ao PROCX), SOMARPRODUTO (SUMPRODUCT), tabelas estruturadas, validação de dados e formatação condicional. O arquivo grava os nomes em inglês; o Excel exibe no idioma do seu programa."),
        ("Como validar", "1) Aba Qualidade: todos PASS. 2) Aba Indicadores, seção 6: todas as conferências OK. 3) Compare os KPIs com scripts/validar_dados.py e docs/TEST_REPORT.md (cálculo independente em Python a partir dos CSVs)."),
        ("Não incluído na V0.1", "Tabelas dinâmicas (PivotTables) não foram criadas nesta versão: os resumos usam fórmulas SOMASES/CONT.SES, que são reconciliáveis com os CSVs. Painel/gráficos: V0.2."),
        ("", ""),
        ("SHA-256 dos CSVs de origem", ""),
    ]
    for nome, h in (meta.get("sha") or {}).items():
        linhas.append((nome, h))
    linhas += [("", ""), ("Uso", "Somente estudo e portfólio. Nenhum uso operacional real. Veja docs/LIMITATIONS.md.")]
    for k, (a, b) in enumerate(linhas, start=4):
        ca, cb = ws.cell(row=k, column=1, value=a or None), ws.cell(row=k, column=2, value=b or None)
        ca.font = f_bold if a else f_norm
        cb.font = f_norm
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
        if a in ("Abas", "SHA-256 dos CSVs de origem"):
            ca.font = f_sec
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120


# ------------------------------------------------------------------ recálculo + cache
class LibreOfficeIndisponivel(RuntimeError):
    """LibreOffice não encontrado: o recálculo (e os testes que dependem dele) não pode rodar."""


# Perfil isolado do LibreOffice com "recalcular sempre ao abrir arquivos OOXML" (OOXMLRecalcMode = 0).
XCU_RECALC = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""
CODIGOS_ERRO = {"#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}


def localizar_soffice():
    """Devolve o caminho do executável do LibreOffice ou None.

    Ordem: 1) variável LOBO_SOFFICE (se definida, é a única fonte: caminho inválido => None,
    para o erro ficar explícito em vez de usar outro programa em silêncio); 2) PATH.
    """
    explicito = os.environ.get("LOBO_SOFFICE")
    if explicito:
        return explicito if Path(explicito).is_file() else None
    for nome in ("soffice", "libreoffice"):
        achado = shutil.which(nome)
        if achado:
            return achado
    return None


def recalcular(caminho, timeout=180):
    """Recalcula (LibreOffice headless) uma CÓPIA; devolve (valores por aba/célula, resumo do recálculo).

    O arquivo original não é modificado. O resumo tem: status, total_errors, error_summary, total_formulas.
    """
    soffice = localizar_soffice()
    if soffice is None:
        raise LibreOfficeIndisponivel(
            "LibreOffice não encontrado. Instale-o e deixe o comando 'soffice' no PATH, "
            "ou defina a variável de ambiente LOBO_SOFFICE com o caminho do executável.")
    with tempfile.TemporaryDirectory(prefix="lobo_recalc_") as pasta:
        pasta = Path(pasta)
        entrada, saida = pasta / "entrada.xlsx", pasta / "saida"
        shutil.copy(caminho, entrada)
        perfil = pasta / "perfil"
        (perfil / "user").mkdir(parents=True)
        (perfil / "user" / "registrymodifications.xcu").write_text(XCU_RECALC, encoding="utf-8")
        cmd = [soffice, "--headless", "--norestore", f"-env:UserInstallation={perfil.as_uri()}",
               "--convert-to", "xlsx", "--outdir", str(saida), str(entrada)]
        try:
            p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise RuntimeError(f"LibreOffice excedeu {timeout}s; nenhuma fórmula foi recalculada.")
        recalculado = saida / "entrada.xlsx"
        if p.returncode != 0 or not recalculado.exists():
            raise RuntimeError(f"LibreOffice falhou (código {p.returncode}): {p.stdout} {p.stderr}")
        wb = load_workbook(recalculado, data_only=True)
        valores = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
                   for ws in wb.worksheets}
        wb.close()
        wf = load_workbook(entrada)
        total = sum(isinstance(c.value, str) and c.value.startswith("=") for ws in wf.worksheets for row in ws.iter_rows() for c in row)
        wf.close()
    erros = Counter(v for aba in valores.values() for v in aba.values() if isinstance(v, str) and v in CODIGOS_ERRO)
    info = {"status": "success", "total_errors": sum(erros.values()), "error_summary": dict(erros), "total_formulas": total}
    return valores, info


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def injetar_cache(cru, valores, saida):
    """Grava os valores calculados como cache (<v>) nas células com fórmula do arquivo cru."""
    with zipfile.ZipFile(cru) as z:
        wbxml = etree.fromstring(z.read("xl/workbook.xml"))
        rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rid2target = {r.get("Id"): r.get("Target") for r in rels}
        RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        aba_xml = {}
        for s in wbxml.find("m:sheets", NS):
            alvo = rid2target[s.get(RID)].lstrip("/")
            aba_xml["xl/" + alvo if not alvo.startswith("xl/") else alvo] = s.get("name")
        with zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as out:
            for item in z.infolist():
                dados = z.read(item.filename)
                if item.filename in aba_xml:
                    vals = valores.get(aba_xml[item.filename], {})
                    raiz = etree.fromstring(dados)
                    for c in raiz.iter("{%s}c" % NS["m"]):
                        f = c.find("m:f", NS)
                        if f is None:
                            continue
                        v = c.find("m:v", NS)
                        if v is None:
                            v = etree.SubElement(c, "{%s}v" % NS["m"])
                        val = vals.get(c.get("r"))
                        if isinstance(val, bool):
                            c.set("t", "b"); v.text = "1" if val else "0"
                        elif isinstance(val, (int, float)):
                            if "t" in c.attrib:
                                del c.attrib["t"]
                            v.text = repr(val)
                        elif isinstance(val, (dt.datetime, dt.date)):
                            if "t" in c.attrib:
                                del c.attrib["t"]
                            v.text = repr(to_excel(val))
                        elif isinstance(val, str) and val.startswith("#"):
                            c.set("t", "e"); v.text = val
                        else:
                            c.set("t", "str"); v.text = "" if val is None else str(val)
                    dados = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
                out.writestr(item, dados)


def sha256(caminho):
    h = hashlib.sha256()
    with open(caminho, "rb") as fh:
        for bloco in iter(lambda: fh.read(65536), b""):
            h.update(bloco)
    return h.hexdigest()


def main():
    from validar_dados import carregar
    from lobo_common import DATA
    dfs = carregar()
    meta = {"sha": {f.name: sha256(f) for f in sorted(DATA.glob("*.csv"))}}
    tmp = Path(tempfile.mkdtemp(prefix="lobo_build_"))
    cru = tmp / "cru.xlsx"
    construir(dfs, cru, meta)
    try:
        valores, info = recalcular(cru)
    except LibreOfficeIndisponivel as e:
        raise SystemExit(f"ERRO: {e}")
    print("recalc:", info)
    if info.get("status") != "success" or info.get("total_errors"):
        raise SystemExit("Erros de fórmula após o recálculo – corrija antes de continuar.")
    EXCEL.mkdir(exist_ok=True)
    destino = EXCEL / NOME_XLSX
    injetar_cache(cru, valores, destino)
    shutil.rmtree(tmp, ignore_errors=True)
    print("gerado:", destino, "sha256:", sha256(destino))


if __name__ == "__main__":
    main()
