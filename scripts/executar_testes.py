# -*- coding: utf-8 -*-
"""Suíte de testes da V0.1. Tudo aqui é EXECUTADO de verdade; o relatório só registra o que rodou.

Uso:
    python executar_testes.py              # roda tudo e (re)escreve docs/TEST_REPORT.md
    python executar_testes.py --verificar  # roda tudo, NÃO escreve o relatório (usado para conferir o ZIP)

Códigos de saída: 0 = tudo executado e PASS; 1 = há FAIL; 2 = nenhum FAIL, mas há testes SKIP
(não executados por falta do LibreOffice; SKIP nunca é contado como PASS).

Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio.
"""
import ast
import csv
import hashlib
import io
import os
import platform
import re
import subprocess
import sys
import tempfile
import tokenize
import zipfile
from collections import Counter
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd

import gerar_dados
import validar_dados
from construir_excel import NOME_XLSX, LibreOfficeIndisponivel, construir, localizar_soffice, recalcular
from lobo_common import (CATEGORIAS, CENTROS, CHECKS, COLS_CONSUMO, COLS_ESTOQUE,
                         COLS_MATERIAIS, DATA, DOCS, ENC, EXCEL, ROOT, SEP, status_estoque)

AVISO = "Dados 100% sintéticos utilizados exclusivamente para estudo e portfólio."
ARQ = {"materiais_ficticios.csv": COLS_MATERIAIS, "estoque_ficticio.csv": COLS_ESTOQUE,
       "consumo_ficticio.csv": COLS_CONSUMO}
XLSX = EXCEL / NOME_XLSX
BOM, CRLF = b"\xef\xbb\xbf", b"\r\n"
_CB = "CB" + "SI"   # termo proibido (montado assim para o próprio script não se autodenunciar)

LINHAS = []            # testes automatizados: (ID, Área, Cenário, Esperado, Obtido, Status)


def t(area, cenario, esperado, obtido, ok=None, pular=None):
    """Registra um teste. `pular` = motivo: o teste NÃO foi executado e vira SKIP (nunca PASS)."""
    if pular:
        LINHAS.append((f"T{len(LINHAS) + 1:02d}", area, cenario, str(esperado), f"NÃO EXECUTADO: {pular}", "SKIP"))
        return None
    ok = (esperado == obtido) if ok is None else ok
    LINHAS.append((f"T{len(LINHAS) + 1:02d}", area, cenario, str(esperado), str(obtido), "PASS" if ok else "FAIL"))
    return ok


def sha256(p):
    h = hashlib.sha256()
    h.update(Path(p).read_bytes())
    return h.hexdigest()


def fmt_int(n):
    return f"{n:,}".replace(",", ".")


def fmt_money(d):
    return f"{d:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ---------------------------------------------------------------- leitura independente (csv puro + Decimal)
def ler_csv(nome, pasta=DATA):
    with open(pasta / nome, encoding=ENC, newline="") as fh:
        return list(csv.DictReader(fh, delimiter=SEP))


def kpis_csv():
    con, est = ler_csv("consumo_ficticio.csv"), ler_csv("estoque_ficticio.csv")
    consumo = [r for r in con if r["Tipo_Movimentacao"] == "Consumo"]
    return {
        "Unidades consumidas": sum(int(r["Quantidade"]) for r in consumo),
        "Custo estimado consumido": sum(Decimal(r["Quantidade"]) * Decimal(r["Custo_Unitario"]) for r in consumo),
        "Movimentações": len(consumo),
        "Materiais distintos movimentados": len({r["Material_ID"] for r in consumo}),
        "Itens abaixo ou iguais ao mínimo": sum(int(r["Estoque_Atual"]) <= int(r["Estoque_Minimo"]) for r in est),
        "Itens críticos abaixo ou iguais ao mínimo": sum(
            r["Criticidade"] == "Alta" and int(r["Estoque_Atual"]) <= int(r["Estoque_Minimo"]) for r in est),
    }


def achar(ws, rotulo, col=1):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == rotulo:
            return r
    raise KeyError(rotulo)


def bloco(ws, titulo, ncols):
    r0 = achar(ws, titulo) + 2
    out, r = [], r0
    while True:
        out.append([ws.cell(r, c).value for c in range(1, ncols + 1)])
        if ws.cell(r, 1).value == "Total":
            return out
        r += 1


# ---------------------------------------------------------------- varredura de PII / segredos
RE_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
RE_FONE = re.compile(r"(?<![\w.,-])(\+?55[\s-]?)?\(?\d{2}\)?[\s-]?9?\d{4}[\s-]\d{4}(?![\w.,-])")
RE_CPF = re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b")
RE_CNPJ = re.compile(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b")
RE_SEGREDO = re.compile(r"(?i)\b(password|passwd|senha|token|api[_-]?key|secret|bearer)\b\s*[:=]\s*\S+")
RE_PREFIXO = re.compile(r"\b(sk-[A-Za-z0-9]{20,}|ghp_[A-Za-z0-9]{20,}|AKIA[0-9A-Z]{16}|xox[baprs]-[A-Za-z0-9-]{10,})")
RE_URL = re.compile(r"(?i)https?://")
RE_HEX = re.compile(r"\b[0-9a-fA-F]{40,}\b")
PALAVRAS_DADOS = re.compile(r"(?i)\b(senha|password|token|api[_-]?key|secret)\b")


def varrer(texto, dados):
    achados = []
    texto = RE_HEX.sub("", texto)
    for nome, rx in (("e-mail", RE_EMAIL), ("telefone", RE_FONE), ("CPF", RE_CPF), ("CNPJ", RE_CNPJ),
                     ("segredo", RE_SEGREDO), ("prefixo de chave", RE_PREFIXO)):
        achados += [f"{nome}: {m.group(0)[:30]}" for m in rx.finditer(texto)]
    if _CB.lower() in texto.lower():
        achados.append("nome de empresa real")
    if dados:
        achados += [f"palavra: {m.group(0)}" for m in PALAVRAS_DADOS.finditer(texto)]
        achados += [f"url: {m.group(0)}" for m in RE_URL.finditer(texto)]
    return achados


def arquivos_texto():
    ext = {".md", ".py", ".csv", ".txt", ".json", ".sha256"}
    return sorted(p for p in ROOT.rglob("*") if p.is_file() and p.suffix in ext and "__pycache__" not in p.parts
                  and p.name != "TEST_REPORT.md")


# ---------------------------------------------------------------- mutação (prova que as verificações DETECTAM erros)
def mutar(dfs):
    mat, est, con = (dfs[k].copy() for k in ("mat", "est", "con"))
    ncon = con["Material_ID"].value_counts().to_dict()
    exp = {c[0]: 0 for c in CHECKS}
    cen = {c[0]: [] for c in CHECKS}

    def add(q, n, desc):
        exp[q] += n
        cen[q].append(desc)

    def setv(df, mid, col, val, key="Material_ID"):
        df.loc[df[key] == mid, col] = val

    # ---- catálogo
    mat = pd.concat([mat, mat[mat.Material_ID == "MAT016"]], ignore_index=True)
    add("Q01", 2, "linha MAT016 duplicada no catálogo")
    setv(mat, "MAT019", "Custo_Unitario", "0")
    add("Q08", 1, "catálogo MAT019 com custo 0"); add("Q20", 1, "estoque MAT019 diverge do catálogo (custo 0)"); add("Q21", ncon["MAT019"], "consumo de MAT019 diverge do catálogo (custo 0)")
    setv(mat, "MAT010", "Fornecedor", "")
    add("Q04", 1, "catálogo MAT010 com Fornecedor vazio"); add("Q36", 1, "Fornecedor vazio também fica fora do padrão fictício"); add("Q23", 1, "estoque MAT010 diverge do catálogo (fornecedor vazio)")
    setv(mat, "MAT018", "Categoria", "Papelaria")
    add("Q29", 1, "catálogo MAT018 com categoria inválida"); add("Q16", 1, "estoque MAT018 diverge do catálogo (categoria)"); add("Q17", ncon["MAT018"], "consumo de MAT018 diverge do catálogo (categoria)")
    setv(mat, "MAT006", "Unidade", "XX")
    add("Q32", 1, "catálogo MAT006 com unidade inválida"); add("Q18", 1, "estoque MAT006 diverge do catálogo (unidade)"); add("Q19", ncon["MAT006"], "consumo de MAT006 diverge do catálogo (unidade)")
    setv(mat, "MAT014", "Criticidade", "Crítica")
    add("Q33", 1, "catálogo MAT014 com criticidade inválida"); add("Q22", 1, "estoque MAT014 diverge do catálogo (criticidade)")
    setv(mat, "MAT009", "Fornecedor", "Distribuidora X")
    add("Q36", 1, "catálogo MAT009 com fornecedor fora do padrão"); add("Q23", 1, "estoque MAT009 diverge do catálogo (fornecedor)")
    # ---- estoque
    est = pd.concat([est, est[est.Material_ID == "MAT016"]], ignore_index=True)
    add("Q02", 2, "linha MAT016 duplicada no estoque")
    setv(est, "MAT011", "Estoque_Atual", "-2"); add("Q26", 1, "estoque MAT011 = -2")
    setv(est, "MAT012", "Estoque_Minimo", "-5"); add("Q27", 1, "mínimo MAT012 = -5")
    setv(est, "MAT013", "Lead_Time_Dias", "0"); setv(est, "MAT017", "Lead_Time_Dias", "2.5")
    add("Q28", 2, "lead time 0 (MAT013) e 2,5 (MAT017)")
    setv(est, "MAT020", "Fornecedor", ""); add("Q05", 1, "estoque MAT020 com Fornecedor vazio"); add("Q23", 1, "estoque MAT020 diverge do catálogo (fornecedor vazio)")
    setv(est, "MAT007", "Custo_Unitario", "-1"); add("Q09", 1, "estoque MAT007 com custo -1"); add("Q20", 1, "estoque MAT007 diverge do catálogo (custo -1)")
    est = pd.concat([est, pd.DataFrame([{"Material_ID": "MAT999", "Material": "Material Fantasma", "Categoria": "EPI", "Unidade": "UN",
                                        "Estoque_Atual": "10", "Estoque_Minimo": "5", "Custo_Unitario": "1.00", "Lead_Time_Dias": "5",
                                        "Fornecedor": "Fornecedor Alfa", "Criticidade": "Baixa"}])], ignore_index=True)
    add("Q11", 1, "estoque com Material_ID inexistente (MAT999)")
    setv(est, "MAT008", "Material", "Disco de Corte 4,5 pol Premium"); add("Q14", 1, "nome do estoque MAT008 alterado")
    setv(est, "MAT009", "Categoria", "EPI"); add("Q16", 1, "categoria do estoque MAT009 = EPI (válida, mas diverge)")
    setv(est, "MAT015", "Unidade", "UN"); add("Q18", 1, "unidade do estoque MAT015 = UN (válida, mas diverge)")
    setv(est, "MAT001", "Criticidade", "Alta"); add("Q22", 1, "criticidade do estoque MAT001 = Alta (diverge)")
    setv(est, "MAT001", "Estoque_Atual", "50"); setv(est, "MAT003", "Estoque_Atual", "13")   # limites do status (sem erro de qualidade)
    # ---- consumo
    limpos = ~con["Material_ID"].isin(["MAT016", "MAT019", "MAT010", "MAT018", "MAT006", "MAT014", "MAT009"])
    pool = con.index[limpos].tolist()
    un = [i for i in pool[16:] if con.loc[i, "Unidade"] == "UN"][0]
    k = pool[:16]; k[12] = un
    con.loc[k[1], "Movimento_ID"] = con.loc[k[0], "Movimento_ID"]; add("Q03", 2, "Movimento_ID repetido em duas linhas")
    con.loc[k[2], "Material_ID"] = "MAT999"; add("Q12", 1, "consumo com MAT999")
    con.loc[k[3], "Quantidade"] = "0"; con.loc[k[4], "Quantidade"] = "-3"; con.loc[k[5], "Quantidade"] = "2.5"
    add("Q07", 3, "quantidades 0, -3 e 2,5")
    con.loc[k[6], "Custo_Unitario"] = "0"; add("Q10", 1, "consumo com custo 0"); add("Q21", 1, "consumo com custo 0 diverge do catálogo")
    con.loc[k[7], "Tipo_Movimentacao"] = ""; add("Q06", 1, "consumo com Tipo vazio"); add("Q35", 1, "Tipo vazio também é tipo inválido")
    con.loc[k[8], "Tipo_Movimentacao"] = "Devolucao"; add("Q35", 1, "tipo 'Devolucao'")
    con.loc[k[9], "Centro_Trabalho"] = "FUNDICAO"; add("Q34", 1, "centro 'FUNDICAO'")
    con.loc[k[10], "Categoria"] = "Papelaria"; add("Q31", 1, "consumo com categoria 'Papelaria'"); add("Q17", 1, "categoria 'Papelaria' no consumo diverge do catálogo")
    con.loc[k[11], "Material"] = "Nome Errado"; add("Q15", 1, "nome do material alterado no consumo")
    con.loc[k[12], "Unidade"] = "PAR"; add("Q19", 1, "unidade do consumo = PAR (válida, mas diverge)")
    con.loc[k[13], "Custo_Unitario"] = "99.99"; add("Q21", 1, "custo do consumo = 99,99 (diverge)")
    con.loc[k[14], "Data"] = "2026-13-45"; add("Q24", 1, "data inexistente 2026-13-45")
    con.loc[k[15], "Data"] = "2026-09-15"; add("Q24", 1, "data fora do período 2026-09-15")
    return {"mat": mat, "est": est, "con": con}, exp, cen


def imports_mortos(caminho):
    """Nomes importados e nunca usados no módulo."""
    tree = ast.parse(caminho.read_text(encoding="utf-8"))
    imp = {}
    for n in ast.walk(tree):
        if isinstance(n, ast.Import):
            imp.update({(a.asname or a.name).split(".")[0]: n.lineno for a in n.names})
        elif isinstance(n, ast.ImportFrom):
            imp.update({(a.asname or a.name): n.lineno for a in n.names})
    usados = {n.id for n in ast.walk(tree) if isinstance(n, ast.Name)}
    return [f"{k} (linha {v})" for k, v in imp.items() if k not in usados]


def fstrings_312(caminho):
    """Barra invertida ou mesma aspa dentro de campo de f-string: só válido em Python >= 3.12 (PEP 701)."""
    if not hasattr(tokenize, "FSTRING_START"):
        return []                       # Python < 3.12: o próprio interpretador já recusaria a sintaxe
    pilha, ruins = [], []
    for tk in tokenize.generate_tokens(io.StringIO(caminho.read_text(encoding="utf-8")).readline):
        if tk.type == tokenize.FSTRING_START:
            pilha.append([tk.string.lstrip("fFrR"), 0])
        elif tk.type == tokenize.FSTRING_END:
            pilha.pop()
        elif pilha and tk.type != tokenize.FSTRING_MIDDLE:
            topo = pilha[-1]
            if tk.type == tokenize.OP and tk.string in "{}":
                topo[1] += 1 if tk.string == "{" else -1
            elif topo[1] > 0 and tk.type == tokenize.STRING and ("\\" in tk.string or tk.string.lstrip("bBrRuU")[0] == topo[0][0]):
                ruins.append(f"linha {tk.start[0]}")
    return ruins


# ---------------------------------------------------------------- testes
def main(escrever):
    dfs = validar_dados.carregar()
    res = validar_dados.executar(dfs)
    tab_q = validar_dados.tabela(res)

    # ===== Linhas
    n_csv = {k: len(v) for k, v in dfs.items()}
    t("Linhas", "materiais_ficticios.csv – nº de linhas de dados", 20, n_csv["mat"])
    t("Linhas", "estoque_ficticio.csv – nº de linhas de dados", 20, n_csv["est"])
    t("Linhas", "consumo_ficticio.csv – nº de movimentações", 180, n_csv["con"])
    wbf = openpyxl.load_workbook(XLSX)
    wbv = openpyxl.load_workbook(XLSX, data_only=True)
    linhas_tb = {n: int(ws.tables[n].ref.split(":")[1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ")) - 1
                 for n, ws in (("tbMateriais", wbf["Materiais"]), ("tbEstoque", wbf["Estoque"]), ("tbConsumo", wbf["Consumo"]))}
    t("Linhas", "Excel: linhas de tbMateriais / tbEstoque / tbConsumo = CSVs", "20/20/180", f"{linhas_tb['tbMateriais']}/{linhas_tb['tbEstoque']}/{linhas_tb['tbConsumo']}")

    # ===== IDs
    t("Unicidade de IDs", "Material_ID único no catálogo (distintos = linhas)", 20, dfs["mat"]["Material_ID"].nunique())
    t("Unicidade de IDs", "Material_ID único no estoque", 20, dfs["est"]["Material_ID"].nunique())
    t("Unicidade de IDs", "Movimento_ID único no consumo", 180, dfs["con"]["Movimento_ID"].nunique())
    t("Unicidade de IDs", "Material_ID = MAT001..MAT020", True, list(dfs["mat"]["Material_ID"]) == [f"MAT{i:03d}" for i in range(1, 21)])
    t("Unicidade de IDs", "Movimento_ID = MOV0001..MOV0180 (sem lacunas)", True, list(dfs["con"]["Movimento_ID"]) == [f"MOV{i:04d}" for i in range(1, 181)])

    # ===== Integridade referencial
    t("Integridade referencial", "consumo -> catálogo: Material_ID órfãos", 0, res["Q12"])
    t("Integridade referencial", "estoque -> catálogo: Material_ID órfãos", 0, res["Q11"])
    t("Integridade referencial", "catálogo -> estoque: materiais sem posição", 0, res["Q13"])
    t("Integridade referencial", "nome/categoria/unidade/custo/criticidade/fornecedor divergentes (Q14–Q23)", 0, sum(res[f"Q{i}"] for i in range(14, 24)))

    # ===== Regras de domínio
    t("Custos", "custos > 0 nas três tabelas (Q08–Q10)", 0, res["Q08"] + res["Q09"] + res["Q10"])
    con_t = ler_csv("consumo_ficticio.csv")
    qs = [int(r["Quantidade"]) for r in con_t]
    t("Quantidades", "quantidade inteira > 0 em todas as movimentações (Q07); mín/máx observados", "0 erros", f"{res['Q07']} erros (mín {min(qs)}, máx {max(qs)})", res["Q07"] == 0 and min(qs) >= 1)
    datas = sorted(r["Data"] for r in con_t)
    t("Datas", "todas as datas válidas e entre 2026-06-01 e 2026-08-31 (Q24)", "0 erros", f"{res['Q24']} erros ({datas[0]} a {datas[-1]})", res["Q24"] == 0)
    meses = sorted({d[:7] for d in datas})
    t("Datas", "três meses consecutivos com movimentação (Q25)", "['2026-06', '2026-07', '2026-08']", str(meses))
    t("Categorias", "categorias válidas nas três tabelas (Q29–Q31)", 0, res["Q29"] + res["Q30"] + res["Q31"])
    cats = {r["Categoria"] for r in ler_csv("materiais_ficticios.csv")}
    t("Categorias", "as 6 categorias do PDF estão representadas no catálogo", sorted(CATEGORIAS), sorted(cats))
    t("Centros", "centros de trabalho na lista permitida (Q34)", 0, res["Q34"])
    usados = {r["Centro_Trabalho"] for r in con_t}
    t("Centros", "os 5 centros permitidos aparecem no consumo", sorted(CENTROS), sorted(usados))
    est_t = ler_csv("estoque_ficticio.csv")
    t("Estoque", "estoque >= 0, mínimo >= 0, lead time inteiro > 0 (Q26–Q28)", 0, res["Q26"] + res["Q27"] + res["Q28"])
    st = [status_estoque(int(r["Estoque_Atual"]), int(r["Estoque_Minimo"])) for r in est_t]
    t("Estoque", "há itens com Atual <= Mínimo (REPOR)", ">= 1", st.count("REPOR"), st.count("REPOR") >= 1)
    t("Estoque", "PDF: pelo menos 5 itens em ATENÇÃO/REPOR", ">= 5", st.count("REPOR") + st.count("ATENÇÃO"), st.count("REPOR") + st.count("ATENÇÃO") >= 5)
    t("Estoque", "há ao menos um item de estoque exatamente igual ao mínimo", ">= 1", sum(int(r["Estoque_Atual"]) == int(r["Estoque_Minimo"]) for r in est_t), any(int(r["Estoque_Atual"]) == int(r["Estoque_Minimo"]) for r in est_t))
    crit = Counter(r["Criticidade"] for r in est_t)
    t("Estoque", "criticidades Baixa, Média e Alta presentes no estoque", "3 níveis", f"{dict(crit)}", set(crit) == {"Baixa", "Média", "Alta"})
    t("Estoque", "pelo menos um item Alta abaixo ou igual ao mínimo (cenário para o KPI 6)", ">= 1", sum(r["Criticidade"] == "Alta" and int(r["Estoque_Atual"]) <= int(r["Estoque_Minimo"]) for r in est_t), True if any(r["Criticidade"] == "Alta" and int(r["Estoque_Atual"]) <= int(r["Estoque_Minimo"]) for r in est_t) else False)
    por_mat = Counter(r["Material_ID"] for r in con_t)
    t("Distribuição", "consumo NÃO é uniforme entre materiais (máx/mín de movimentações >= 3)", ">= 3", f"{max(por_mat.values()) / min(por_mat.values()):.1f}", max(por_mat.values()) / min(por_mat.values()) >= 3)
    por_mes = Counter(r["Data"][:7] for r in con_t)
    t("Distribuição", "consumo NÃO é uniforme entre meses (movimentações por mês distintas)", ">= 2 valores distintos", str(dict(sorted(por_mes.items()))), len(set(por_mes.values())) >= 2)

    # ===== Formato CSV
    for nome, cols in ARQ.items():
        bruto = (DATA / nome).read_bytes()
        texto = bruto.decode("utf-8")
        linhas = list(csv.reader(texto.lstrip("\ufeff").splitlines(), delimiter=SEP))
        t("Formato CSV", f"{nome}: cabeçalho exato", cols, linhas[0])
        t("Formato CSV", f"{nome}: separador ';' e nº de campos constante", {len(cols)}, {len(l) for l in linhas})
        t("Formato CSV", f"{nome}: UTF-8 com BOM, finais de linha LF", "BOM=True, CRLF=False", f"BOM={bruto.startswith(BOM)}, CRLF={CRLF in bruto}", bruto.startswith(BOM) and CRLF not in bruto)
    padrao = re.compile(r"^\d+\.\d{2}$")
    t("Formato CSV", "Custo_Unitario com ponto decimal e 2 casas nos 3 CSVs", 0,
      sum(not padrao.match(r["Custo_Unitario"]) for nome in ARQ for r in ler_csv(nome) if "Custo_Unitario" in r))

    # ===== Reprodutibilidade
    with tempfile.TemporaryDirectory() as tmp:
        gerar_dados.main(Path(tmp))
        iguais = all(sha256(Path(tmp) / n) == sha256(DATA / n) for n in ARQ)
    t("Reprodutibilidade", "gerar_dados.py (mesma semente) reproduz os 3 CSVs byte a byte", True, iguais)

    # ===== Excel: estrutura
    t("Excel – estrutura", "abas na ordem", ["LEIA_ME", "Materiais", "Estoque", "Consumo", "Qualidade", "Indicadores", "Parametros"], wbf.sheetnames)
    tabs = {n: ws.tables[n].ref for ws in wbf.worksheets for n in ws.tables}
    t("Excel – estrutura", "tabelas estruturadas e intervalos", {"tbMateriais": "A1:G21", "tbEstoque": "A1:L21", "tbConsumo": "A1:L181", "tbQualidade": "A9:E45"}, tabs)
    nomes = sorted(wbf.defined_names.keys())
    t("Excel – estrutura", "nomes definidos (listas e parâmetros)", sorted(["lstCategorias", "lstUnidades", "lstCriticidade", "lstCentros", "lstTipoMov", "lstMeses", "pMargemAtencao", "pInicioPeriodo", "pFimPeriodo"]), nomes)
    with zipfile.ZipFile(XLSX) as z:
        partes = z.namelist()
        wbxml = z.read("xl/workbook.xml").decode("utf-8")
    t("Excel – estrutura", "sem vínculos externos (externalLinks)", 0, sum(p.startswith("xl/externalLinks") for p in partes))
    t("Excel – estrutura", "recálculo total ao abrir (fullCalcOnLoad)", True, "fullCalcOnLoad" in wbxml)
    dv = {ws.title: len(ws.data_validations.dataValidation) for ws in wbf.worksheets}
    t("Excel – estrutura", "validações de dados por aba (Materiais/Estoque/Consumo)", "4/6/8", f"{dv['Materiais']}/{dv['Estoque']}/{dv['Consumo']}")
    cf = {ws.title: len(ws.conditional_formatting) for ws in wbf.worksheets}
    t("Excel – estrutura", "intervalos com formatação condicional (Materiais/Estoque/Qualidade/Indicadores)", "1/3/1/3", f"{cf['Materiais']}/{cf['Estoque']}/{cf['Qualidade']}/{cf['Indicadores']}")

    # ===== Excel: fórmulas
    formulas, sem_cache, erros = 0, [], []
    inventario = Counter()
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    for fn in re.findall(r"([A-Z]+)\(", c.value):
                        inventario[fn] += 1
                    v = wbv[ws.title][c.coordinate].value
                    if v is None:
                        sem_cache.append(f"{ws.title}!{c.coordinate}")
                    elif isinstance(v, str) and v.startswith("#"):
                        erros.append(f"{ws.title}!{c.coordinate}")
    t("Excel – fórmulas", "fórmulas reais no arquivo (>= 800)", ">= 800", formulas, formulas >= 800)
    t("Excel – fórmulas", "células de fórmula sem valor em cache", 0, len(sem_cache))
    t("Excel – fórmulas", "células com erro (#VALUE!, #NAME?, #REF!…) nos valores calculados", 0, len(erros))
    esperadas = ["SUMIFS", "COUNTIFS", "COUNTIF", "IF", "IFERROR", "INDEX", "MATCH", "SUMPRODUCT", "SUM", "ROUND", "COUNTBLANK"]
    t("Excel – fórmulas", "funções do PDF presentes (SOMASES, CONT.SES, SE, SEERRO, ÍNDICE/CORRESP…)", "todas > 0", {f: inventario[f] for f in esperadas}, all(inventario[f] > 0 for f in esperadas))
    constantes = [c.coordinate for row in wbf["Indicadores"].iter_rows() for c in row if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    t("Excel – fórmulas", "aba Indicadores: números digitados à mão (constantes numéricas)", 0, len(constantes))
    ws_c, ws_e = wbv["Consumo"], wbv["Estoque"]
    rows_c = ler_csv("consumo_ficticio.csv")
    ok_cm = ok_mes = 0
    for i, r in enumerate(rows_c, start=2):
        ok_cm += abs(ws_c.cell(i, 11).value - float(Decimal(r["Quantidade"]) * Decimal(r["Custo_Unitario"]))) < 1e-9
        ok_mes += ws_c.cell(i, 12).value == r["Data"][:7]
    t("Excel – fórmulas", "Custo_Movimentado = Quantidade x Custo_Unitario nas 180 linhas", 180, ok_cm)
    t("Excel – fórmulas", "Mes = aaaa-mm da Data nas 180 linhas", 180, ok_mes)
    ok_st = ok_cr = 0
    for i, r in enumerate(est_t, start=2):
        s = status_estoque(int(r["Estoque_Atual"]), int(r["Estoque_Minimo"]))
        ok_st += ws_e.cell(i, 11).value == s
        ok_cr += ws_e.cell(i, 12).value == ("Sim" if r["Criticidade"] == "Alta" and s == "REPOR" else "Não")
    t("Excel – fórmulas", "Status_Estoque (Excel) = regra em Python nas 20 linhas", 20, ok_st)
    t("Excel – fórmulas", "Critico_Em_Risco (Excel) = regra em Python nas 20 linhas", 20, ok_cr)
    wq = wbv["Qualidade"]
    vet_x = {CHECKS[k][0]: wq.cell(10 + k, 4).value for k in range(len(CHECKS))}
    t("Excel – fórmulas", "aba Qualidade: vetor de erros do Excel (Q01–Q36) = vetor do Python", res, vet_x)
    t("Excel – fórmulas", "aba Qualidade: 36 testes, 36 PASS, 0 FAIL", "36/36/0", f"{wq['B5'].value}/{wq['B6'].value}/{wq['B7'].value}")
    wi = wbv["Indicadores"]
    t("Excel – fórmulas", "aba Indicadores: conferências internas", "12 de 12", wi.cell(achar(wi, "Conferências internas OK"), 2).value)

    # resumos vs. cálculo independente
    est_map = {r["Material_ID"]: r for r in est_t}
    def agrupa(chave):
        d = {}
        for r in rows_c:
            x = d.setdefault(chave(r), [0, 0, Decimal(0)])
            x[0] += 1; x[1] += int(r["Quantidade"]); x[2] += Decimal(r["Quantidade"]) * Decimal(r["Custo_Unitario"])
        return d
    def confere(titulo, chaves, chave, ncols=5):
        ref, bl, bad = agrupa(chave), bloco(wi, titulo, ncols), 0
        for k, (lab, mov, un, custo) in zip(chaves, [b[:4] for b in bl[:-1]]):
            g = ref.get(k, [0, 0, Decimal(0)])
            bad += not (lab == k and mov == g[0] and un == g[1] and abs(custo - float(g[2])) < 0.005)
        return bad
    t("Excel – resumos", "resumo por mês = agrupamento independente dos CSVs", 0, confere("2. Resumo por mês", ["2026-06", "2026-07", "2026-08"], lambda r: r["Data"][:7]))
    t("Excel – resumos", "resumo por categoria = agrupamento independente", 0, confere("3. Resumo por categoria", CATEGORIAS, lambda r: r["Categoria"]))
    t("Excel – resumos", "resumo por centro de trabalho = agrupamento independente", 0, confere("4. Resumo por centro de trabalho", CENTROS, lambda r: r["Centro_Trabalho"]))
    t("Excel – resumos", "resumo por material = agrupamento independente (20 materiais)", 0, confere("5. Resumo por material", [f"MAT{i:03d}" for i in range(1, 21)], lambda r: r["Material_ID"]))

    # ===== Reconciliação dos 6 KPIs
    kp = kpis_csv()
    recon = []
    for nome, esperado in kp.items():
        r = achar(wi, nome)
        x = wi.cell(r, 2).value
        ok = abs(Decimal(str(x)) - Decimal(esperado)) < Decimal("0.005")
        recon.append((nome, esperado, f"{x:.2f}" if isinstance(esperado, Decimal) else x, ok))
        t("Reconciliação de KPIs", f"KPI '{nome}': Excel = cálculo independente dos CSVs", f"{esperado}", f"{x}", ok)

    # ===== Detecção (mutação): as verificações realmente pegam erros?
    dfm, exp, cen = mutar(dfs)
    py_m = validar_dados.executar(dfm)
    valores, info, sem_lo = None, None, None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "mutado.xlsx"
            construir(dfm, arq)
            valores, info = recalcular(arq)
    except LibreOfficeIndisponivel as e:
        sem_lo = str(e)
    ne = len(dfm["est"])
    est_p = []
    for _, r in dfm["est"].iterrows():
        s = status_estoque(float(r["Estoque_Atual"]), float(r["Estoque_Minimo"]))
        est_p.append((s, "Sim" if r["Criticidade"] == "Alta" and s == "REPOR" else "Não"))
    if valores is not None:
        qx = valores["Qualidade"]
        xl_m = {CHECKS[k][0]: qx.get(f"D{10 + k}") for k in range(len(CHECKS))}
    else:
        xl_m = {c[0]: "NÃO EXECUTADO" for c in CHECKS}
    det = []
    for cid, teste, _, _ in CHECKS:
        if valores is None:
            estado = "SKIP" if exp[cid] == py_m[cid] else "FAIL"      # só o lado Python foi executado
        else:
            estado = "PASS" if exp[cid] == py_m[cid] == xl_m[cid] else "FAIL"
        det.append((cid, "; ".join(cen[cid]) or "(nenhum erro injetado)", exp[cid], py_m[cid], xl_m[cid], estado))
    t("Detecção de erros", "dataset com erros injetados: Python detecta exatamente o esperado nos 36 testes", "36/36", f"{sum(exp[c[0]] == py_m[c[0]] for c in CHECKS)}/36", all(exp[c[0]] == py_m[c[0]] for c in CHECKS))
    if valores is None:
        t("Detecção de erros", "dataset com erros injetados: fórmulas do Excel detectam exatamente o esperado nos 36 testes", "36/36", "", pular=sem_lo)
        t("Detecção de erros", "Excel marca FAIL exatamente nos testes com erro injetado", sum(v > 0 for v in exp.values()), "", pular=sem_lo)
        t("Detecção de erros", "Status_Estoque/Critico_Em_Risco no Excel = Python, incluindo estoque negativo e limites (Atual = Mín, = Mín x 1,25, logo acima)", f"{ne}/{ne}", "", pular=sem_lo)
        t("Detecção de erros", "limites: MAT002 (30/30) = REPOR; MAT001 (50 = 40 x 1,25) = ATENÇÃO; MAT003 (13 > 12,5) = OK", "REPOR/ATENÇÃO/OK", "", pular=sem_lo)
    else:
        t("Detecção de erros", "dataset com erros injetados: fórmulas do Excel detectam exatamente o esperado nos 36 testes", "36/36", f"{sum(exp[c[0]] == xl_m[c[0]] for c in CHECKS)}/36", all(exp[c[0]] == xl_m[c[0]] for c in CHECKS))
        fails_x = [qx.get(f"E{10 + k}") for k in range(len(CHECKS))]
        t("Detecção de erros", "Excel marca FAIL exatamente nos testes com erro injetado", sum(v > 0 for v in exp.values()), fails_x.count("FAIL"))
        est_x = [(valores["Estoque"].get(f"K{i}"), valores["Estoque"].get(f"L{i}")) for i in range(2, ne + 2)]
        t("Detecção de erros", "Status_Estoque/Critico_Em_Risco no Excel = Python, incluindo estoque negativo e limites (Atual = Mín, = Mín x 1,25, logo acima)", f"{ne}/{ne}", f"{sum(a == b for a, b in zip(est_x, est_p))}/{ne}", est_x == est_p)
        lim = {mid: valores["Estoque"].get(f"K{2 + list(dfm['est'].Material_ID).index(mid)}") for mid in ("MAT002", "MAT001", "MAT003")}
        t("Detecção de erros", "limites: MAT002 (30/30) = REPOR; MAT001 (50 = 40 x 1,25) = ATENÇÃO; MAT003 (13 > 12,5) = OK", "REPOR/ATENÇÃO/OK", f"{lim['MAT002']}/{lim['MAT001']}/{lim['MAT003']}")

    # ===== Segurança
    achados = []
    for p in arquivos_texto():
        achados += [f"{p.relative_to(ROOT)} -> {a}" for a in varrer(p.read_text(encoding="utf-8-sig"), p.suffix == ".csv")]
    txt_xlsx = "\n".join(str(c.value) for ws in wbf.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    achados += [f"XLSX -> {a}" for a in varrer(txt_xlsx, True)]
    t("Segurança", "varredura de PII/credenciais em CSVs, XLSX, scripts e documentos (e-mail, telefone, CPF/CNPJ, senha/token/chave, nome de empresa real)", 0, len(achados))
    proibidos = [p.name for p in ROOT.rglob("*") if p.is_file() and (p.name in (".env", ".DS_Store", "Thumbs.db") or p.suffix in (".pbix", ".pem", ".key") or p.name.startswith("credentials"))]
    t("Segurança", "sem .env, .pbix, chaves ou credenciais no pacote", 0, len(proibidos))
    t("Segurança", "fornecedores fictícios ('Fornecedor …') (Q36)", 0, res["Q36"])

    # ===== Documentação
    exigidos = ["README.md", "CHANGELOG.md", "docs/DATA_DICTIONARY.md", "docs/LIMITATIONS.md", "docs/V0_1_NOTES.md",
                "data/materiais_ficticios.csv", "data/estoque_ficticio.csv", "data/consumo_ficticio.csv", f"excel/{NOME_XLSX}"]
    t("Documentação", "arquivos obrigatórios existem (o ZIP e o TEST_REPORT são conferidos à parte)", [], [e for e in exigidos if not (ROOT / e).exists()])
    pastas = ["prompts", "powerbi", "metrics", "n8n", "images", "data", "excel", "docs"]
    t("Documentação", "estrutura de pastas do enunciado", [], [d for d in pastas if not (ROOT / d).is_dir()])
    docs = {n: (ROOT / n).read_text(encoding="utf-8") for n in exigidos if n.endswith(".md")}
    t("Documentação", "aviso 'Dados 100% sintéticos…' em README, DATA_DICTIONARY, LIMITATIONS, V0_1_NOTES", [], [n for n, x in docs.items() if n != "CHANGELOG.md" and AVISO not in x])
    t("Documentação", "CHANGELOG inicia com 'V0.1 — Foundation'", True, "V0.1 — Foundation" in docs["CHANGELOG.md"])
    dd = docs["docs/DATA_DICTIONARY.md"]
    linhas_dd = [[c.strip() for c in l.strip().strip("|").split("|")] for l in dd.splitlines() if l.startswith("|") and not set(l) <= set("|-: ")]
    faltam, exemplo_falso = [], []
    for nome, cols in ARQ.items():
        stem = nome.replace(".csv", "")
        rows = ler_csv(nome)
        for c in cols:
            achou = [l for l in linhas_dd if len(l) >= 6 and l[0].strip("`") == stem and l[1].strip("`") == c]
            if not achou:
                faltam.append(f"{stem}.{c}")
            elif achou[0][5].strip("`") not in {r[c] for r in rows}:
                exemplo_falso.append(f"{stem}.{c}='{achou[0][5]}'")
    t("Documentação", "DATA_DICTIONARY descreve todas as colunas dos 3 CSVs", [], faltam)
    t("Documentação", "exemplos do DATA_DICTIONARY existem de fato nas colunas dos CSVs", [], exemplo_falso)
    calc_ok = all(any(l[1].strip("`") == c for l in linhas_dd) for c in ("Status_Estoque", "Critico_Em_Risco", "Custo_Movimentado", "Mes"))
    t("Documentação", "DATA_DICTIONARY inclui as 4 colunas calculadas do Excel", True, calc_ok)
    faltam_kpi = []
    for nome, val in kp.items():
        v = fmt_money(val) if isinstance(val, Decimal) else fmt_int(val)
        for arq in ("README.md", "docs/V0_1_NOTES.md"):
            if not re.search(rf"\|\s*{re.escape(nome)}\s*\|\s*{re.escape(v)}\s*\|", docs[arq]):
                faltam_kpi.append(f"{arq}:{nome}={v}")
    t("Documentação", "README e V0_1_NOTES citam os 6 KPIs com os valores reais (tabela)", [], faltam_kpi)
    lim_txt = docs["docs/LIMITATIONS.md"].lower()
    termos = ["sintétic", "três meses", "custos fictícios", "snapshot", "compras", "contratos", "previsão", "uso operacional"]
    t("Documentação", "LIMITATIONS cobre os 8 pontos exigidos", [], [x for x in termos if x not in lim_txt])
    notas = docs["docs/V0_1_NOTES.md"]
    t("Documentação", "V0_1_NOTES contém 'O que o autor precisa saber explicar' e os 4 blocos (O QUE FOI FEITO / POR QUE FOI FEITO / COMO VALIDAR / O QUE PRECISO SABER EXPLICAR)", True,
      "O que o autor precisa saber explicar" in notas and all(b in notas for b in ("O QUE FOI FEITO", "POR QUE FOI FEITO", "COMO VALIDAR", "O QUE PRECISO SABER EXPLICAR")))
    proib = ["skl" + "earn", "stats" + "models", "prop" + "het", "Random" + "Forest", "ARI" + "MA", "FORE" + "CAST", "TR" + "END(", "GRO" + "WTH("]
    achou_scripts = [w for p in (ROOT / "scripts").glob("*.py") for w in proib if w.lower() in p.read_text(encoding="utf-8").lower() and p.name != "executar_testes.py"]
    formulas_txt = " ".join(c.value for ws in wbf.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    achou_xlsx = [w for w in proib if w in formulas_txt.upper()]
    t("Escopo", "Projeto 1 sem forecasting: nenhuma biblioteca/modelo de previsão nos scripts e nenhuma função FORECAST/TREND/GROWTH nas fórmulas do Excel", [], achou_scripts + achou_xlsx)
    # controle positivo: o detector de PII/segredos realmente detecta?
    amostras = {"e-mail": "contato: joao" + "@" + "exemplo.com", "telefone": "ligar (24) 9" + "9999-1234", "CPF": "cpf 123.456." + "789-09",
                "CNPJ": "cnpj 12.345.678" + "/0001-95", "segredo": "sen" + "ha = abc123", "chave": "api" + "_key: XYZ", "empresa real": "fornecedor " + _CB}
    nao_detectou = [k for k, v in amostras.items() if not varrer(v, False)]
    t("Segurança", "controle positivo: o detector de PII/segredos sinaliza 7 amostras plantadas (e-mail, telefone, CPF, CNPJ, senha, chave, empresa real)", [], nao_detectou)
    t("Segurança", "controle negativo: texto limpo do catálogo não gera alerta", 0, len(varrer("Disco de Corte 7 pol; Fornecedor Alfa; 12.50; 2026-06-03", True)))

    # ===== V0.1.1 – portabilidade e regressão contra a V0.1 congelada (T91 em diante; T01–T90 permanecem idênticos aos da V0.1)
    amb = ["/m" + "nt/", "/ho" + "me/", "/t" + "mp/", "C:" + "\\", "recal" + "c.py", "LOBO_" + "RECALC"]
    scripts = sorted((ROOT / "scripts").glob("*.py"))
    t("Portabilidade", "scripts sem caminhos absolutos do ambiente original e sem referência a script de recálculo externo ao projeto", [], [f"{p.name}: {a}" for p in scripts for a in amb if a in p.read_text(encoding="utf-8")])
    ant = os.environ.get("LOBO_SOFFICE")
    os.environ["LOBO_SOFFICE"] = str(ROOT / "nao_existe" / "soffice")
    try:
        achado = localizar_soffice()
        try:
            recalcular(XLSX)
            erro = "nenhum erro"
        except LibreOfficeIndisponivel as e:
            erro = "LibreOfficeIndisponivel" if "LOBO_SOFFICE" in str(e) else "mensagem sem orientação"
    finally:
        if ant is None:
            os.environ.pop("LOBO_SOFFICE", None)
        else:
            os.environ["LOBO_SOFFICE"] = ant
    t("Portabilidade", "LOBO_SOFFICE inválido: localizar_soffice() = None e recalcular() levanta LibreOfficeIndisponivel com orientação (sem travar)", "None / LibreOfficeIndisponivel", f"{achado} / {erro}")
    lo = localizar_soffice()
    t("Portabilidade", "LibreOffice localizado neste ambiente (variável LOBO_SOFFICE ou PATH)", "encontrado", "encontrado" if lo else "não encontrado", pular=None if lo else "LibreOffice ausente (restrição do ambiente, não do projeto)")
    if lo:
        with tempfile.TemporaryDirectory() as tmp:
            cru = Path(tmp) / "cru.xlsx"
            construir(dfs, cru, {"sha": {f.name: sha256(f) for f in sorted(DATA.glob("*.csv"))}})
            novo_val, novo_info = recalcular(cru)
            wf_novo = openpyxl.load_workbook(cru)
            n_cel = dif_f = dif_v = 0
            for ws in wbf.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if c.value is None:
                            continue
                        n_cel += 1
                        dif_f += wf_novo[ws.title][c.coordinate].value != c.value
                        if isinstance(c.value, str) and c.value.startswith("="):
                            a, b = wbv[ws.title][c.coordinate].value, novo_val[ws.title].get(c.coordinate)
                            dif_v += not ((abs(a - b) < 1e-9) if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) else a == b)
            wf_novo.close()
        t("Portabilidade", "reconstruir o XLSX a partir dos CSVs (recálculo portátil) reproduz o XLSX entregue: fórmulas/constantes e valores calculados de todas as células", "0 / 0 diferenças", f"{dif_f} / {dif_v} diferenças em {n_cel} células ({novo_info['total_formulas']} fórmulas, {novo_info['total_errors']} erros)", dif_f == 0 and dif_v == 0 and novo_info["total_errors"] == 0)
    else:
        t("Portabilidade", "reconstruir o XLSX a partir dos CSVs (recálculo portátil) reproduz o XLSX entregue", "0 / 0 diferenças", "", pular="LibreOffice ausente (restrição do ambiente, não do projeto)")
    with tempfile.TemporaryDirectory() as outro:
        r = subprocess.run([sys.executable, str(ROOT / "scripts" / "validar_dados.py")], cwd=outro, capture_output=True, text=True,
                           encoding="utf-8", env={**os.environ, "PYTHONIOENCODING": "utf-8"})
    t("Portabilidade", "validar_dados.py executado a partir de outro diretório de trabalho: código de saída 0 e 36 PASS", "0 / 36 PASS / 0 FAIL", f"{r.returncode} / " + ("36 PASS / 0 FAIL" if "36 PASS; 0 FAIL" in r.stdout else "resultado inesperado"))
    base_txt = (DOCS / "V0_1_BASELINE_SHA256.txt").read_text(encoding="utf-8")
    base = [l.split("  ", 1) for l in base_txt.splitlines() if l.strip() and not l.startswith("#")]
    hist = {"docs/TEST_REPORT.md": "docs/TEST_REPORT_V0_1.md"}      # o relatório da V0.1 foi preservado com outro nome
    t("Regressão vs V0.1", "arquivos que a V0.1.1 não pode alterar (CSVs, XLSX, regras em lobo_common/gerar_dados/validar_dados, DATA_DICTIONARY, V0_1_NOTES, relatório da V0.1) seguem byte a byte iguais à base congelada",
      [], [rel for h, rel in base if sha256(ROOT / hist.get(rel, rel)) != h])
    t("Regressão vs V0.1", "a base congelada registra exatamente 10 arquivos protegidos", 10, len(base))
    novos_arq = ["docs/V0_1_1_NOTES.md", "docs/TEST_REPORT_V0_1.md", "docs/V0_1_BASELINE_SHA256.txt", "docs/manual/inspecao_visual.md", "docs/manual/observacoes_execucao.md", "requirements.txt"]
    t("Documentação V0.1.1", "arquivos novos da V0.1.1 presentes", [], [a for a in novos_arq if not (ROOT / a).exists()])
    ch = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    nt = (DOCS / "V0_1_1_NOTES.md").read_text(encoding="utf-8")
    t("Documentação V0.1.1", "CHANGELOG registra 'V0.1.1 — Portable Baseline' acima da seção V0.1 (preservada); V0_1_1_NOTES traz o aviso de dados sintéticos", True,
      "V0.1.1 — Portable Baseline" in ch and "V0.1 — Foundation" in ch and ch.index("V0.1.1 — Portable Baseline") < ch.index("V0.1 — Foundation") and AVISO in nt)
    req = (ROOT / "requirements.txt").read_text(encoding="utf-8").lower()
    t("Portabilidade", "requirements.txt declara pandas, openpyxl e lxml", [], [x for x in ("pandas", "openpyxl", "lxml") if x not in req])
    rel_manual = [a for a in ("docs/manual/inspecao_visual.md", "docs/manual/observacoes_execucao.md") if (ROOT / a).exists()]
    t("Portabilidade", "o gerador do relatório lê as notas manuais de dentro do pacote (não de pastas externas)", 2, len(rel_manual))
    t("Qualidade do código", "sem imports não usados nos scripts", [], sum(([f"{p.name}: {m}" for m in imports_mortos(p)] for p in scripts), []))
    t("Qualidade do código", "sem f-strings que exijam Python >= 3.12 (barra invertida ou mesma aspa dentro do campo)", [], sum(([f"{p.name}: {m}" for m in fstrings_312(p)] for p in scripts), []))

    # ===== saída
    total = len(LINHAS)
    falhas = [l for l in LINHAS if l[5] == "FAIL"]
    pulados = [l for l in LINHAS if l[5] == "SKIP"]
    det_fail = [d for d in det if d[5] == "FAIL"]
    det_skip = [d for d in det if d[5] == "SKIP"]
    q_fail = int((tab_q["Status"] == "FAIL").sum())
    print(f"Testes automatizados: {total} | PASS {total - len(falhas) - len(pulados)} | FAIL {len(falhas)} | SKIP {len(pulados)}")
    print(f"Verificações de qualidade (Python): {len(tab_q)} | FAIL {q_fail}")
    print(f"Detecção por mutação: {len(det)} | FAIL {len(det_fail)} | SKIP {len(det_skip)}")
    for l in falhas:
        print("FAIL:", l)
    for l in pulados:
        print("SKIP:", l[0], l[2][:90], "->", l[4])
    for d in det_fail:
        print("FAIL DET:", d)
    if escrever:
        (DOCS / "TEST_REPORT.md").write_text(relatorio(tab_q, det, recon, dfs, info), encoding="utf-8")
        print("TEST_REPORT.md escrito")
    if falhas or det_fail or q_fail:
        return 1
    return 2 if (pulados or det_skip) else 0


def esc(s):
    return str(s).replace("|", "\\|").replace("\n", " ")


def relatorio(tab_q, det, recon, dfs, info_mut):
    def linhas_md(cab, rows):
        out = ["| " + " | ".join(cab) + " |", "|" + "|".join("---" for _ in cab) + "|"]
        out += ["| " + " | ".join(esc(c) for c in r) + " |" for r in rows]
        return "\n".join(out)
    exe = localizar_soffice()
    try:
        lo = subprocess.run([exe, "--version"], capture_output=True, text=True, timeout=60).stdout.strip() if exe else "LibreOffice NÃO encontrado"
    except Exception:
        lo = "LibreOffice (versão não obtida)"
    total = len(LINHAS)
    npass = sum(l[5] == "PASS" for l in LINHAS)
    nskip = sum(l[5] == "SKIP" for l in LINHAS)
    manual = DOCS / "manual" / "inspecao_visual.md"
    insp = manual.read_text(encoding="utf-8") if manual.exists() else "Inspeção visual não registrada."
    obs = DOCS / "manual" / "observacoes_execucao.md"
    observ = obs.read_text(encoding="utf-8") if obs.exists() else "Nenhuma observação registrada."
    parts = [
        "# TEST_REPORT — Lobo Insights Industrial V0.1.1 (Portable Baseline)",
        f"> {AVISO}",
        "",
        "Este relatório é **gerado por `scripts/executar_testes.py`**: cada linha PASS/FAIL abaixo foi produzida por uma execução real. "
        "Nada foi marcado PASS sem executar; teste que não pôde rodar aparece como **SKIP** (nunca PASS). Para reproduzir: `python scripts/executar_testes.py --verificar`. "
        "O relatório da V0.1 original está preservado, sem alteração, em `docs/TEST_REPORT_V0_1.md`.",
        "",
        "## 1. Resumo",
        f"- Testes automatizados: **{total}** — PASS **{npass}**, FAIL **{sum(l[5] == 'FAIL' for l in LINHAS)}**, SKIP **{nskip}** (T01–T90 são os mesmos testes da V0.1; T91 em diante são novos da V0.1.1).",
        f"- Verificações de qualidade dos CSVs (Python, Q01–Q36): **{len(tab_q)}** — PASS **{int((tab_q['Status'] == 'PASS').sum())}**, FAIL **{int((tab_q['Status'] == 'FAIL').sum())}**.",
        f"- Testes de detecção por injeção de erros: **{len(det)}** — PASS **{sum(d[5] == 'PASS' for d in det)}**, FAIL **{sum(d[5] == 'FAIL' for d in det)}**, SKIP **{sum(d[5] == 'SKIP' for d in det)}**.",
        f"- Ambiente: Python {platform.python_version()}, pandas {pd.__version__}, openpyxl {openpyxl.__version__}, {lo}, {platform.system()} {platform.machine()}.",
        f"- SHA-256 do XLSX (`excel/{NOME_XLSX}`): `{sha256(XLSX)}`",
        "- SHA-256 dos CSVs: " + "; ".join(f"`{n}` = `{sha256(DATA / n)}`" for n in ARQ),
        "",
        "## 2. TESTADO × NÃO TESTADO",
        "**TESTADO (executado nesta V0.1):** geração determinística dos CSVs; 36 verificações de qualidade em Python e o mesmo conjunto em fórmulas do Excel (recalculadas no LibreOffice); "
        "detecção de erros injetados (Python e Excel); estrutura do XLSX (abas, tabelas, nomes, validações, formatação condicional, ausência de vínculos externos); "
        "colunas calculadas; resumos por mês/categoria/centro/material; reconciliação dos 6 KPIs com um cálculo independente (módulo `csv` + `Decimal`); "
        "formato/codificação dos CSVs; varredura por regex de PII/credenciais; consistência da documentação com os dados. "
        "**Adicionados na V0.1.1:** ausência de caminhos absolutos do ambiente original; comportamento controlado sem LibreOffice; reconstrução do XLSX idêntica à entregue; execução a partir de outro diretório; "
        "regressão byte a byte dos arquivos protegidos contra a base V0.1 congelada; imports não usados; sintaxe compatível com Python anterior à 3.12.",
        "",
        "**NÃO TESTADO (declarado, não assumido):**",
        "- Abertura em Microsoft Excel real (Windows/Mac/Online). O arquivo foi lido com openpyxl e recalculado/renderizado com LibreOffice; comportamento no Excel não foi verificado.",
        "- Excel com idioma pt-BR: as fórmulas são gravadas com nomes em inglês e o Excel deve traduzi-las ao abrir; isso não foi verificado.",
        "- Semântica de célula vazia em `COUNTIF` (usada nos testes Q35 e derivados) foi verificada no LibreOffice; o Excel real pode diferir em casos extremos.",
        "- Diferença de maiúsculas/minúsculas: `COUNTIFS` do Excel não diferencia caixa; o validador Python diferencia. Divergências só de caixa não foram testadas no Excel.",
        "- Preenchimento automático das colunas calculadas ao adicionar linhas na tabela do Excel (não configurado nem testado).",
        "- Tabelas dinâmicas (PivotTables): **não foram criadas** nesta versão.",
        "- Dashboard/gráficos (V0.2), IA executiva (V0.3), classificador (V0.4), avaliação (V0.5), Power BI (V0.6), n8n (V0.7), GitHub (V1.0): fora do escopo, nada foi testado.",
        "- Testes do PDF sobre IA e automação (número inventado, prioridade urgente, JSON inválido, entrada vazia): pertencem a versões futuras.",
        "- A varredura de PII é heurística por regex: não prova a ausência absoluta de dado sensível; complementa a revisão humana.",
        "- **V0.1.1:** execução em outros ambientes. Só foram verificados Linux x86_64 e Python 3.12.3 (com as versões de bibliotecas e do LibreOffice acima). Windows, macOS, outras versões de Python/pandas/openpyxl/LibreOffice **não foram testados**.",
        "- **V0.1.1:** instalação limpa por `pip install -r requirements.txt` em ambiente virtual novo (sem acesso à rede neste ambiente); as bibliotecas usadas foram as já instaladas.",
        "- **V0.1.1:** uso de `LOBO_SOFFICE` apontando para um `soffice` válido em local não padrão e codificação do console do Windows nos textos impressos: não testados.",
        "- O conteúdo do ZIP final é conferido após o empacotamento, fora deste relatório (um arquivo não pode conter a prova do próprio ZIP).",
        "",
        "## 3. Verificações de qualidade dos dados (Python, sobre os CSVs)",
        "Resultado da execução de `scripts/validar_dados.py`. As mesmas 36 verificações existem como fórmulas na aba **Qualidade** do Excel.",
        "",
        linhas_md(["Teste", "Regra", "Resultado", "Quantidade de erros", "Status"],
                  [(f"{r.ID} · {r.Teste}", r.Regra, r.Resultado, r["Quantidade de erros"], r.Status) for _, r in tab_q.iterrows()]),
        "",
        "## 4. Testes automatizados",
        linhas_md(["ID", "Área", "Cenário", "Esperado", "Obtido", "PASS/FAIL"], LINHAS),
        "",
        "## 5. Reconciliação dos 6 KPIs (Excel × CSVs)",
        "O valor do CSV vem de `csv.DictReader` + `Decimal` (sem pandas e sem as fórmulas do Excel); o valor do Excel é o número calculado da aba Indicadores.",
        "",
        linhas_md(["KPI", "Calculado dos CSVs", "Excel (aba Indicadores)", "Status"], [(n, e, x, "PASS" if ok else "FAIL") for n, e, x, ok in recon]),
        "",
        "## 6. Testes de detecção: as verificações realmente pegam erros?",
        "Um dataset **com erros injetados de propósito** (cópia em memória; os CSVs oficiais não são alterados) é validado em Python e em um XLSX temporário com as mesmas fórmulas. "
        "\"Esperado\" foi definido à mão a partir do que foi injetado (incluindo efeitos colaterais, ex.: custo 0 no catálogo também diverge do estoque e do consumo).",
        "",
        linhas_md(["Teste", "Erros injetados", "Esperado", "Python", "Excel", "PASS/FAIL"], det),
        "",
        "## 7. Inspeção visual do Excel",
        insp,
        "",
        "## 8. Registro de problemas e observações da execução",
        observ,
        "",
    ]
    return "\n".join(parts)


if __name__ == "__main__":
    sys.exit(main(escrever="--verificar" not in sys.argv))
