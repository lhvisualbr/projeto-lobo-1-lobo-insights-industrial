import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"

if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))


import construir_excel_v0_2_1 as gen  # noqa: E402
import lobo_release as rel  # noqa: E402
from construir_excel import localizar_soffice  # noqa: E402
from lobo_common import EXCEL  # noqa: E402
from lobo_locale import textos_localizados_equivalentes  # noqa: E402


class TestXlsxLocalePortability(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not localizar_soffice():
            raise unittest.SkipTest(
                "LibreOffice não localizado neste ambiente."
            )

        cls._tmp = tempfile.TemporaryDirectory(
            prefix="lobo_v030_locale_"
        )

        cls.original = EXCEL / rel.NOME_XLSX
        cls.reconstruido = (
            Path(cls._tmp.name) / rel.NOME_XLSX
        )

        if not cls.original.is_file():
            raise AssertionError(
                f"XLSX oficial não encontrado: {cls.original}"
            )

        cls.info = gen.gerar(cls.reconstruido)

        cls.formulas_original = openpyxl.load_workbook(
            cls.original,
            data_only=False,
        )

        cls.formulas_reconstruido = openpyxl.load_workbook(
            cls.reconstruido,
            data_only=False,
        )

        cls.valores_original = openpyxl.load_workbook(
            cls.original,
            data_only=True,
        )

        cls.valores_reconstruido = openpyxl.load_workbook(
            cls.reconstruido,
            data_only=True,
        )

    @classmethod
    def tearDownClass(cls):
        if hasattr(cls, "_tmp"):
            cls._tmp.cleanup()

    def test_recalculo_libreoffice_sem_erros(self):
        self.assertEqual(
            self.info.get("status"),
            "success",
            self.info,
        )

        self.assertFalse(
            self.info.get("total_errors"),
            self.info,
        )

    def test_mesmas_abas(self):
        self.assertEqual(
            self.formulas_original.sheetnames,
            self.formulas_reconstruido.sheetnames,
        )

    def test_formulas_e_constantes_sao_identicas(self):
        diferencas = []

        for ws_original in self.formulas_original.worksheets:
            ws_novo = self.formulas_reconstruido[
                ws_original.title
            ]

            max_row = max(
                ws_original.max_row,
                ws_novo.max_row,
            )

            max_col = max(
                ws_original.max_column,
                ws_novo.max_column,
            )

            for row in range(1, max_row + 1):
                for col in range(1, max_col + 1):
                    a = ws_original.cell(row, col).value
                    b = ws_novo.cell(row, col).value

                    if a != b:
                        diferencas.append(
                            (
                                ws_original.title,
                                ws_original.cell(
                                    row,
                                    col,
                                ).coordinate,
                                a,
                                b,
                            )
                        )

        self.assertEqual(
            diferencas,
            [],
            "Diferenças reais de fórmula/constante: "
            + repr(diferencas[:20]),
        )

    def test_valores_em_cache_aceitam_apenas_localidade(self):
        diferencas_reais = []
        diferencas_localidade = []

        for ws_original in self.valores_original.worksheets:
            ws_novo = self.valores_reconstruido[
                ws_original.title
            ]

            max_row = max(
                ws_original.max_row,
                ws_novo.max_row,
            )

            max_col = max(
                ws_original.max_column,
                ws_novo.max_column,
            )

            for row in range(1, max_row + 1):
                for col in range(1, max_col + 1):
                    celula_original = ws_original.cell(
                        row,
                        col,
                    )

                    celula_nova = ws_novo.cell(
                        row,
                        col,
                    )

                    a = celula_original.value
                    b = celula_nova.value

                    if a == b:
                        continue

                    if (
                        isinstance(a, (int, float))
                        and isinstance(b, (int, float))
                        and not isinstance(a, bool)
                        and not isinstance(b, bool)
                    ):
                        if abs(a - b) <= 1e-9:
                            continue

                    if textos_localizados_equivalentes(
                        a,
                        b,
                    ):
                        diferencas_localidade.append(
                            (
                                ws_original.title,
                                celula_original.coordinate,
                                a,
                                b,
                            )
                        )
                        continue

                    diferencas_reais.append(
                        (
                            ws_original.title,
                            celula_original.coordinate,
                            a,
                            b,
                        )
                    )

        self.assertEqual(
            diferencas_reais,
            [],
            "Diferenças reais de valor encontradas: "
            + repr(diferencas_reais[:20]),
        )

        print(
            "\nDiferenças aceitas somente por localidade:",
            len(diferencas_localidade),
        )

        for item in diferencas_localidade:
            print(
                f"  {item[0]}!{item[1]}: "
                f"{item[2]!r} -> {item[3]!r}"
            )

    def test_quantidade_de_graficos_preservada(self):
        graficos_original = sum(
            len(ws._charts)
            for ws in self.formulas_original.worksheets
        )

        graficos_reconstruido = sum(
            len(ws._charts)
            for ws in self.formulas_reconstruido.worksheets
        )

        self.assertEqual(
            graficos_original,
            graficos_reconstruido,
        )

        self.assertEqual(
            graficos_original,
            9,
        )


if __name__ == "__main__":
    unittest.main()